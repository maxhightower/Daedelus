"""V1.2: native Office adapters (xlsx / docx / pptx) and cross-artifact dependencies.

These run against the real libraries (openpyxl, python-docx, python-pptx) and, where present,
LibreOffice for recalculation, previews and reopen checks. Tests that need LibreOffice skip
with the reason when it is missing (CI installs it).
"""

from __future__ import annotations

import os
import shutil
import zipfile
from pathlib import Path

import pytest

from daedelus import dependencies as dp
from daedelus.adapters import AdapterError, get_adapter
from daedelus.adapters.office import common as oc
from daedelus.artifacts import create_artifact, native_path
from daedelus.editing import EditRejected, apply_manual_edit
from daedelus.models import PlannedOperation as P
from daedelus.store import Workspace

# CI sets DAEDELUS_REQUIRE_LIBREOFFICE=1: these tests must then run (and fail), never skip
LO = pytest.mark.skipif(oc.soffice() is None and not os.environ.get("DAEDELUS_REQUIRE_LIBREOFFICE"),
                        reason="LibreOffice (soffice) not installed")


@pytest.fixture()
def st(tmp_path):
    ws = Workspace(tmp_path / "ws")
    _, store = ws.create_project("office", "")
    yield store
    ws.close()


def _book(st):
    a, _ = create_artifact(st, name="Book", adapter="spreadsheet", template="data", params={
        "sheets": [{"id": "data", "title": "Data", "rows": [
            ["Decade", "Mean"], ["1960s", 319.5], ["1970s", 330.1], ["1980s", 345.5]]},
            {"id": "summary", "title": "Summary"}]})
    apply_manual_edit(st, a.id, [
        P(op="set_cells", component_id="summary", params={"cells": {
            "A1": "Decade", "B1": "Mean", "A2": "=Data!A2", "B2": "=Data!B2", "A3": "=Data!A3",
            "B3": "=Data!B3", "A4": "=Data!A4", "B4": "=Data!B4", "C4": "=B4-B2"}}),
        P(op="define_name", params={"id": "summary_values", "sheet": "summary", "ref": "A1:B4",
                                    "name": "SummaryValues"}),
        P(op="add_chart", params={"id": "decade_chart", "sheet": "summary", "type": "column",
                                  "data": "B1:B4", "categories": "A2:A4", "anchor": "E2",
                                  "title": "Decadal mean"})], message="build")
    return st.get_artifact(a.id)


def _report(st):
    d, _ = create_artifact(st, name="Report", adapter="document", template="blank",
                           params={"title": "Report"})
    apply_manual_edit(st, d.id, [
        P(op="add_heading", params={"id": "results", "text": "Results"}),
        P(op="add_paragraph", params={"id": "results_text", "text": "(pending)",
                                      "parent": "results"}),
        P(op="add_heading", params={"id": "discussion", "text": "Discussion"}),
        P(op="add_paragraph", params={"id": "disc1", "parent": "discussion",
                                      "text": "Concentrations keep rising."}),
        P(op="add_paragraph", params={"id": "disc2", "parent": "discussion",
                                      "text": "Growth has accelerated."})], message="doc")
    return st.get_artifact(d.id)


def _deck(st):
    p, _ = create_artifact(st, name="Deck", adapter="presentation", template="blank", params={})
    apply_manual_edit(st, p.id, [
        P(op="add_slide", params={"id": "data_slide", "layout": "title_only", "title": "Data"}),
        P(op="add_slide", params={"id": "findings", "layout": "title_content",
                                  "title": "Findings"})], message="deck")
    return st.get_artifact(p.id)


def _props(st, a):
    a = st.get_artifact(a.id)
    return get_adapter(a.adapter).inspect(native_path(st, a), a.entry)


# ------------------------------------------------------------------ spreadsheet
def test_workbook_hierarchy_formulas_and_identity(st):
    a = _book(st)
    kinds = {c.id: c.kind for c in a.components}
    assert kinds["data"] == kinds["summary"] == "sheet"
    assert kinds["summary_values"] == "range" and kinds["decade_chart"] == "chart"
    import openpyxl
    wb = openpyxl.load_workbook(native_path(st, a) / a.entry)
    # formulas are stored as formulas, never flattened to values
    assert wb["Summary"]["B4"].value == "=Data!B4" and wb["Summary"]["C4"].value == "=B4-B2"
    assert "SummaryValues" in wb.defined_names
    # stable ids survive a reopen (sidecar + native locators)
    ins = _props(st, a)
    assert {"data", "summary", "summary_values", "decade_chart"} <= set(ins.properties)


@LO
def test_formula_values_are_recalculated_on_a_copy(st):
    a = _book(st)
    ad = get_adapter("spreadsheet")
    vals = ad.values(native_path(st, a), a.entry, "Summary!A1:C4")
    assert vals[3][1] == pytest.approx(345.5) and vals[3][2] == pytest.approx(26.0)
    import openpyxl
    wb = openpyxl.load_workbook(native_path(st, a) / a.entry)
    assert wb["Summary"]["C4"].value == "=B4-B2"  # native file still holds the formula
    rep = ad.validate(native_path(st, a), a.entry, ["file_reopens"], {})
    names = {c.name: c.passed for c in rep.checks}
    assert names.get("formulas_preserved") and names.get("recalculated")


def test_rename_sheet_rewrites_references(st):
    a = _book(st)
    apply_manual_edit(st, a.id, [P(op="rename_sheet", component_id="data",
                                   params={"title": "Observations"})], message="rename")
    import openpyxl
    a = st.get_artifact(a.id)
    wb = openpyxl.load_workbook(native_path(st, a) / a.entry)
    assert wb["Summary"]["B4"].value == "=Observations!B4"
    assert next(c for c in a.components if c.id == "data").name == "Observations"


def test_insert_rows_refused_when_references_would_break(st):
    a = _book(st)
    with pytest.raises(EditRejected, match="insert refused"):
        apply_manual_edit(st, a.id, [P(op="insert_rows", component_id="data",
                                       params={"index": 3, "amount": 1})], message="insert")
    # rollback: the workbook is unchanged
    import openpyxl
    a = st.get_artifact(a.id)
    assert openpyxl.load_workbook(native_path(st, a) / a.entry)["Data"]["A3"].value == "1970s"


def test_edit_outside_target_range_is_rejected(st):
    a = _book(st)
    with pytest.raises(EditRejected, match="outside the target"):
        apply_manual_edit(st, a.id, [P(op="set_cells", component_id="summary_values",
                                       params={"cells": {"D9": 1}})], message="bad")


def test_missing_target_is_reported(st):
    a = _book(st)
    with pytest.raises(EditRejected, match="unknown component|not found"):
        apply_manual_edit(st, a.id, [P(op="set_cells", component_id="no_such_range",
                                       params={"cells": {"A1": 1}})], message="bad")


def test_macro_workbook_is_refused_not_destroyed(st, tmp_path):
    a = _book(st)
    path = native_path(st, a) / a.entry
    before = path.read_bytes()
    with zipfile.ZipFile(path, "a") as z:  # simulate a VBA project in the package
        z.writestr("xl/vbaProject.bin", b"\x00fake-vba")
    tampered = path.read_bytes()
    with pytest.raises(EditRejected, match="(?i)macro|vba|unsupported"):
        apply_manual_edit(st, a.id, [P(op="set_cells", component_id="data",
                                       params={"cells": {"B2": 1.0}})], message="edit")
    assert path.read_bytes() == tampered != before  # file left exactly as found


def test_chart_survives_unrelated_edits(st):
    a = _book(st)
    apply_manual_edit(st, a.id, [P(op="set_cells", component_id="data",
                                   params={"cells": {"B2": 320.0}})], message="edit")
    a = st.get_artifact(a.id)
    with zipfile.ZipFile(native_path(st, a) / a.entry) as z:
        assert any(n.startswith("xl/charts/chart") for n in z.namelist())


# ------------------------------------------------------------------ document
def test_document_sections_and_scoped_text_edit(st):
    d = _report(st)
    kinds = {c.id: (c.kind, c.parent_id) for c in d.components}
    assert kinds["results"][0] == "section" and kinds["results_text"] == ("paragraph", "results")
    before = _props(st, d).properties
    apply_manual_edit(st, d.id, [P(op="set_text", component_id="disc2",
                                   params={"text": "Growth accelerated after 1990."})],
                      message="edit")
    after = _props(st, d).properties
    assert after["disc2"]["text"] == "Growth accelerated after 1990."
    for cid in ("results_text", "disc1", "results", "discussion"):
        assert after[cid]["text"] == before[cid]["text"]


def test_document_keeps_content_added_outside_daedelus(st):
    d = _report(st)
    import docx
    path = native_path(st, d) / d.entry
    doc = docx.Document(str(path))
    doc.add_paragraph("A note typed in Word.")
    doc.save(str(path))
    apply_manual_edit(st, d.id, [P(op="set_text", component_id="disc1",
                                   params={"text": "Concentrations rose."})], message="edit")
    texts = [p.text for p in docx.Document(str(path)).paragraphs]
    assert "A note typed in Word." in texts and "Concentrations rose." in texts


def test_document_table_and_replace_section(st):
    d = _report(st)
    apply_manual_edit(st, d.id, [
        P(op="add_table", params={"id": "t1", "after": "results_text",
                                  "rows": [["a", "b"], ["1", "2"]], "header": True}),
        P(op="replace_section", component_id="discussion",
          params={"paragraphs": ["Only this paragraph remains."]})], message="edit")
    props = _props(st, d).properties
    assert props["t1"]["rows"][1] == ["1", "2"]
    import docx
    a = st.get_artifact(d.id)
    texts = [p.text for p in docx.Document(str(native_path(st, a) / a.entry)).paragraphs]
    assert "Only this paragraph remains." in texts and "Growth has accelerated." not in texts
    assert "(pending)" in texts  # other sections untouched


# ------------------------------------------------------------------ presentation
def test_presentation_slides_native_text_and_chart(st):
    p = _deck(st)
    apply_manual_edit(st, p.id, [
        P(op="set_text", component_id="findings.body", params={"paragraphs": ["One", "Two"]}),
        P(op="add_chart", params={"id": "c1", "slide": "data_slide", "type": "column",
                                  "box": [0.1, 0.25, 0.8, 0.6],
                                  "categories": ["x", "y"], "series": {"s": [1, 2]}})],
        message="edit")
    from pptx import Presentation
    a = st.get_artifact(p.id)
    prs = Presentation(str(native_path(st, a) / a.entry))
    shapes = [s for s in prs.slides[0].shapes]
    assert any(s.has_chart for s in shapes)  # a native chart object, not a picture
    body = [s for s in prs.slides[1].shapes if s.has_text_frame and "One" in s.text_frame.text]
    assert body and [pp.text for pp in body[0].text_frame.paragraphs] == ["One", "Two"]
    apply_manual_edit(st, p.id, [P(op="update_chart_data", component_id="c1", params={
        "categories": ["x", "y", "z"], "series": {"s": [3, 4, 5]}})],
        message="data")
    a = st.get_artifact(p.id)
    prs = Presentation(str(native_path(st, a) / a.entry))
    ch = next(s for s in prs.slides[0].shapes if s.has_chart).chart
    assert list(ch.plots[0].series[0].values) == [3, 4, 5]


def test_presentation_slide_move_keeps_ids(st):
    p = _deck(st)
    apply_manual_edit(st, p.id, [P(op="move_slide", component_id="findings",
                                   params={"index": 0})], message="move")
    a = st.get_artifact(p.id)
    slides = [c.id for c in a.components if c.kind == "slide"]
    from pptx import Presentation
    prs = Presentation(str(native_path(st, a) / a.entry))
    assert prs.slides[0].shapes.title.text == "Findings"
    assert set(slides) == {"data_slide", "findings"}


# ------------------------------------------------------------------ dependencies
def _links(st, x, d, p):
    E = dp.End
    return [
        dp.add_dependency(st, E(artifact_id=x.id, component_id="summary_values"),
                          E(artifact_id=d.id, component_id="results_table"), target_kind="table",
                          options={"create": {"after": "results_text"}}),
        dp.add_dependency(st, E(artifact_id=x.id, component_id="summary_values"),
                          E(artifact_id=d.id, component_id="results_text"),
                          options={"template": "Mean rose from {B2:.1f} in the {A2} to "
                                               "{B4:.1f} in the {A4}."}),
        dp.add_dependency(st, E(artifact_id=x.id, component_id="summary_values"),
                          E(artifact_id=p.id, component_id="deck_chart"), target_kind="chart",
                          options={"create": {"slide": "data_slide",
                                              "box": [0.1, 0.25, 0.8, 0.65]}}),
        dp.add_dependency(st, E(artifact_id=d.id, component_id="discussion"),
                          E(artifact_id=p.id, component_id="findings.body"),
                          options={"max_paragraphs": 3}),
    ]


@LO
def test_dependencies_sync_and_selective_update(st):
    x, d, p = _book(st), _report(st), _deck(st)
    deps = _links(st, x, d, p)
    assert {s["status"] for s in dp.all_status(st)} == {"never_synced"}
    rep = dp.sync(st)
    assert not rep["failed"] and len(rep["revisions"]) == 2  # one revision per target artifact
    assert {s["status"] for s in dp.all_status(st)} == {"synced"}
    assert _props(st, d).properties["results_text"]["text"] == \
        "Mean rose from 319.5 in the 1960s to 345.5 in the 1980s."

    apply_manual_edit(st, x.id, [P(op="set_cells", component_id="data",
                                   params={"cells": {"B4": 350.0}})], message="data change")
    stat = {s["target"]["component_id"]: s["status"] for s in dp.all_status(st)}
    assert stat == {"results_table": "stale", "results_text": "stale", "deck_chart": "stale",
                    "findings.body": "synced"}
    disc_before = _props(st, p).properties["findings.body"]
    rep = dp.sync(st)
    changed = {r["artifact_id"]: set(r["changed"]) for r in rep["revisions"]}
    assert changed[d.id] == {"results_table", "results_text"}
    assert changed[p.id] == {"deck_chart"}
    assert _props(st, p).properties["findings.body"] == disc_before
    assert "350.0" in _props(st, d).properties["results_text"]["text"]
    assert {s["status"] for s in dp.all_status(st)} == {"synced"}
    assert len(deps) == 4


def test_dependency_validation(st):
    x, d = _book(st), _report(st)
    E = dp.End
    with pytest.raises(ValueError, match="no dependency route"):
        dp.add_dependency(st, E(artifact_id=d.id, component_id="disc1"),
                          E(artifact_id=x.id, component_id="summary_values"))
    with pytest.raises(ValueError, match="not in"):
        dp.add_dependency(st, E(artifact_id=x.id, component_id="nope"),
                          E(artifact_id=d.id, component_id="results_text"))
    with pytest.raises(ValueError, match="options.create"):
        dp.add_dependency(st, E(artifact_id=x.id, component_id="summary_values"),
                          E(artifact_id=d.id, component_id="new_table"), target_kind="table")
    dp.add_dependency(st, E(artifact_id=x.id, component_id="summary_values"),
                      E(artifact_id=d.id, component_id="results_text"))
    p = _deck(st)
    dp.add_dependency(st, E(artifact_id=d.id, component_id="discussion"),
                      E(artifact_id=p.id, component_id="findings.body"))
    # p -> x would close a cycle x -> d -> p -> x (needs a route from a deck component; use
    # a document paragraph route back to the workbook is not offered, so test via doc->book)
    with pytest.raises(ValueError, match="cycle|route"):
        dp.add_dependency(st, E(artifact_id=d.id, component_id="disc1"),
                          E(artifact_id=x.id, component_id="summary_values"))


def test_missing_source_component_is_reported(st):
    x, d = _book(st), _report(st)
    dep = dp.add_dependency(st, dp.End(artifact_id=x.id, component_id="summary_values"),
                            dp.End(artifact_id=d.id, component_id="results_text"))
    apply_manual_edit(st, x.id, [P(op="delete_component", component_id="summary_values")],
                      message="delete")
    s = dp.status(st, dep)
    assert s["status"] == "missing_source"
    rep = dp.sync(st)
    assert not rep["revisions"]


# ------------------------------------------------------------------ recipes / pipeline
@LO
def test_research_pipeline_recipe_builds_native_files(st):
    from daedelus import ingest
    from daedelus.engine import Engine
    from daedelus.models import (SourceBinding, TargetSelector, Workflow, WorkflowEdge,
                                 WorkflowNode)
    assets = Path(__file__).resolve().parents[1] / "daedelus" / "demo_assets" / "v12"
    data = ingest.register_file(st, assets / "co2_annmean_mlo.csv", name="CO2")
    notes = ingest.register_file(st, assets / "research_notes.md", name="Notes")
    for s in (data, notes):
        ingest.ingest(st, st.get_source(s.id))
    wb, _ = create_artifact(st, name="Analysis", adapter="spreadsheet", template="blank")
    doc, _ = create_artifact(st, name="Report", adapter="document", template="blank",
                             params={"title": "CO2"})
    for s, a, role in ((data, wb, "reference"), (notes, doc, "reference")):
        st.save_binding(SourceBinding(source_id=s.id, role=role,
                                      target=TargetSelector(scope="artifact", artifact_id=a.id)))
    nodes = [WorkflowNode(id="src", type="sources", label="S", config={}),
             WorkflowNode(id="aw", type="artifact", label="W", config={"artifact_id": wb.id}),
             WorkflowNode(id="ad", type="artifact", label="D", config={"artifact_id": doc.id}),
             WorkflowNode(id="w", type="agent", label="w", config={
                 "instructions": "Build the analysis workbook.", "fan_out": False}),
             WorkflowNode(id="d", type="agent", label="d", config={
                 "instructions": "Write the report.", "fan_out": False}),
             WorkflowNode(id="deps", type="dependencies", label="deps", config={"links": [
                 {"source": {"artifact_id": wb.id, "component_id": "summary_values"},
                  "target": {"artifact_id": doc.id, "component_id": "results_table"},
                  "target_kind": "table", "options": {"create": {"parent": "results"}}}]})]
    edges = [WorkflowEdge(id="1", source="aw", source_port="artifact", target="w",
                          target_port="artifact"),
             WorkflowEdge(id="2", source="ad", source_port="artifact", target="d",
                          target_port="artifact"),
             WorkflowEdge(id="3", source="src", source_port="sources", target="w",
                          target_port="sources"),
             WorkflowEdge(id="4", source="src", source_port="sources", target="d",
                          target_port="sources"),
             WorkflowEdge(id="5", source="w", source_port="revision", target="deps",
                          target_port="after"),
             WorkflowEdge(id="6", source="d", source_port="revision", target="deps",
                          target_port="after")]
    wf = Workflow(name="p", nodes=nodes, edges=edges)
    st.save_workflow(wf)
    ex = Engine(st).execute(wf.id)
    assert ex.status.value == "succeeded", ex.error
    w = st.get_artifact(wb.id)
    assert {"data_table", "dataset", "summary_values", "trend_chart", "decade_chart"} <= \
        {c.id for c in w.components}
    vals = get_adapter("spreadsheet").values(native_path(st, w), w.entry, "Summary!A1:B9")
    assert vals[1][0] == "1950s" and 300 < vals[1][1] < 330  # computed by AVERAGEIFS
    plan = ex.run("w").units[0].plan
    assert plan["deterministic"] and "recipe" in plan["notes"][0]
    rows = _props(st, st.get_artifact(doc.id)).properties["results_table"]["rows"]
    assert rows[0][0] == "Decade" and rows[1][0] == "1950s"


def test_truncated_extract_is_flagged(tmp_path):
    big = tmp_path / "big.csv"
    big.write_text("year,v\n" + "".join(f"{1000 + i},{i}\n" for i in range(6000)))
    from daedelus import office
    sheet = office.read_spreadsheet(big)["sheets"][0]
    # the analysis-workbook recipe refuses truncated extracts instead of analysing a prefix
    assert sheet["truncated"] and sheet["row_count"] == 6001 and len(sheet["rows"]) == 5000


@LO
def test_previews_are_rendered_by_libreoffice(st):
    p = _deck(st)
    rev = st.get_revision(st.get_artifact(p.id).head_revision_id)
    assert rev.previews.get("pdf") and rev.previews.get("render")
    assert st.abs(rev.previews["render"]).stat().st_size > 1000


def test_unsupported_feature_detection_lists_parts(tmp_path):
    src = tmp_path / "x.pptx"
    from pptx import Presentation
    Presentation().save(str(src))
    with zipfile.ZipFile(src, "a") as z:
        z.writestr("ppt/embeddings/oleObject1.bin", b"x")
    assert any("embedded" in f.lower() or "ole" in f.lower()
               for f in oc.unsupported_features(src, "pptx"))
    shutil.rmtree(tmp_path, ignore_errors=True)
    with pytest.raises(AdapterError):
        get_adapter("spreadsheet").create(tmp_path / "n", "nope", {})
