"""Connector contracts against synthetic HTTP fixtures (mock transport; no live account).

The fixture payloads follow the documented response shapes of Microsoft Graph drive items
and Google Drive v3 files. They are hand-written, not recorded: they verify request
construction, version/conflict handling, export selection and the security rules - not that
a live tenant behaves the same.
"""

from __future__ import annotations

import os
import io
import json

import httpx
import pytest

from daedelus import connectors as cn
from daedelus.adapters.office import common as oc
from daedelus.store import Workspace

# CI sets DAEDELUS_REQUIRE_LIBREOFFICE=1: these tests must then run (and fail), never skip
LO = pytest.mark.skipif(oc.soffice() is None and not os.environ.get("DAEDELUS_REQUIRE_LIBREOFFICE"),
                        reason="LibreOffice (soffice) not installed")


def _docx_bytes(text: str) -> bytes:
    import docx
    d = docx.Document()
    d.add_paragraph(text)
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


class FakeGraph:
    def __init__(self):
        self.etag = '"{A1}",1'
        self.content = _docx_bytes("Hello from OneDrive")
        self.calls: list[httpx.Request] = []

    def __call__(self, req: httpx.Request) -> httpx.Response:
        self.calls.append(req)
        url = str(req.url)
        if req.url.host == "contoso.sharepoint.com":
            assert "authorization" not in req.headers  # token never forwarded to content host
            return httpx.Response(200, content=self.content)
        assert req.headers["authorization"] == "Bearer T"
        if req.method == "GET" and url.startswith(
                "https://graph.microsoft.com/v1.0/me/drive/items/ITEM1?"):
            return httpx.Response(200, json={
                "id": "ITEM1", "name": "Report.docx", "eTag": self.etag, "size": 1234,
                "lastModifiedDateTime": "2026-10-01T10:00:00Z", "webUrl": "https://x/Report",
                "file": {"mimeType": cn.OOXML["docx"]}})
        if req.method == "GET" and url.endswith("/items/ITEM1/content"):
            return httpx.Response(302, headers={
                "location": "https://contoso.sharepoint.com/download?tempauth=abc"})
        if req.method == "PUT" and url.endswith("/items/ITEM1/content"):
            if req.headers.get("if-match") != self.etag:
                return httpx.Response(412, json={"error": {"code": "resourceModified"}})
            self.content = req.content
            self.etag = '"{A1}",2'
            return httpx.Response(200, json={"id": "ITEM1", "name": "Report.docx",
                                             "eTag": self.etag,
                                             "file": {"mimeType": cn.OOXML["docx"]}})
        if url.endswith("/items/EVIL/content"):
            return httpx.Response(302, headers={"location": "https://attacker.example/x"})
        if "/items/EVIL" in url:
            return httpx.Response(200, json={"id": "EVIL", "name": "x.docx", "eTag": "e",
                                             "file": {"mimeType": cn.OOXML["docx"]}})
        return httpx.Response(404)


class FakeDrive:
    def __init__(self):
        self.version = "7"
        self.calls: list[httpx.Request] = []

    def __call__(self, req: httpx.Request) -> httpx.Response:
        self.calls.append(req)
        url = req.url
        if url.path == "/drive/v3/files/GDOC" and "alt" not in url.params:
            return httpx.Response(200, json={
                "id": "GDOC", "name": "Plan", "mimeType": "application/vnd.google-apps.document",
                "version": "3", "modifiedTime": "2026-10-01T10:00:00Z"})
        if url.path == "/drive/v3/files/GDOC/export":
            assert url.params["mimeType"] == cn.OOXML["docx"]
            return httpx.Response(200, content=_docx_bytes("Exported from Google Docs"))
        if url.path == "/drive/v3/files/XL1" and url.params.get("alt") == "media":
            import openpyxl
            wb = openpyxl.Workbook()
            wb.active["A1"] = 42
            b = io.BytesIO()
            wb.save(b)
            return httpx.Response(200, content=b.getvalue())
        if url.path == "/drive/v3/files/XL1":
            return httpx.Response(200, json={
                "id": "XL1", "name": "Data.xlsx", "mimeType": cn.OOXML["xlsx"],
                "version": self.version, "size": "4096"})
        if url.path == "/upload/drive/v3/files/XL1" and req.method == "PATCH":
            assert url.params["uploadType"] == "media"
            self.version = str(int(self.version) + 1)
            return httpx.Response(200, json={"id": "XL1", "name": "Data.xlsx",
                                             "mimeType": cn.OOXML["xlsx"],
                                             "version": self.version})
        return httpx.Response(404)


@pytest.fixture()
def st(tmp_path):
    ws = Workspace(tmp_path / "ws")
    _, store = ws.create_project("conn", "")
    yield store
    ws.close()


def test_status_reports_blocked_without_credentials(monkeypatch):
    monkeypatch.delenv("DAEDELUS_MSGRAPH_TOKEN", raising=False)
    monkeypatch.delenv("DAEDELUS_GOOGLE_TOKEN", raising=False)
    reg = cn.registry()
    for name in ("msgraph", "google"):
        s = reg[name].status()
        assert not s.available and s.reason.startswith("blocked")
        with pytest.raises(cn.ConnectorBlocked):
            reg[name].get("ABC")


def test_graph_import_edit_publish_and_conflict(st):
    fake = FakeGraph()
    g = cn.GraphConnector(httpx.MockTransport(fake), token="T")
    art, rev = cn.import_remote(st, g, "ITEM1")
    assert art.adapter == "document" and art.metadata["remote"]["version"] == '"{A1}",1'
    from daedelus.adapters import get_adapter
    from daedelus.artifacts import native_path
    ins = get_adapter("document").inspect(native_path(st, art), art.entry)
    assert any(p.get("text") == "Hello from OneDrive" for p in ins.properties.values())
    # publish with the imported eTag succeeds and records the new version
    new = cn.publish_remote(st, art.id, g)
    assert new.version == '"{A1}",2'
    assert st.get_artifact(art.id).metadata["remote"]["version"] == '"{A1}",2'
    # someone edits the remote copy -> our stored eTag is stale -> 412 -> conflict, no overwrite
    fake.etag = '"{A1}",3'
    before = fake.content
    with pytest.raises(cn.ConnectorConflict):
        cn.publish_remote(st, art.id, g)
    assert fake.content == before


def test_graph_refuses_foreign_redirect_and_bad_ids(tmp_path):
    g = cn.GraphConnector(httpx.MockTransport(FakeGraph()), token="T")
    with pytest.raises(cn.ConnectorError, match="refused redirect"):
        g.download("EVIL", tmp_path)
    for bad in ("../me", "a/b", "x?y=1", "", "a" * 300):
        with pytest.raises(cn.ConnectorError, match="invalid remote id"):
            g.get(bad)


def test_google_native_doc_is_exported_and_not_overwritten(st):
    fake = FakeDrive()
    gd = cn.GoogleDriveConnector(httpx.MockTransport(fake), token="T")
    art, _ = cn.import_remote(st, gd, "GDOC")
    remote = art.metadata["remote"]
    assert remote["exported"] and remote["format"] == "docx" and art.adapter == "document"
    with pytest.raises(cn.ConnectorError, match="Google-native"):
        cn.publish_remote(st, art.id, gd)
    assert not any(r.method == "PATCH" for r in fake.calls)


def test_google_binary_publish_checks_version(st):
    fake = FakeDrive()
    gd = cn.GoogleDriveConnector(httpx.MockTransport(fake), token="T")
    art, _ = cn.import_remote(st, gd, "XL1")
    assert art.adapter == "spreadsheet" and art.metadata["remote"]["version"] == "7"
    assert cn.publish_remote(st, art.id, gd).version == "8"
    fake.version = "11"  # changed remotely
    with pytest.raises(cn.ConnectorConflict):
        cn.publish_remote(st, art.id, gd)


def test_only_allowlisted_hosts():
    g = cn.GraphConnector(httpx.MockTransport(FakeGraph()), token="T")
    assert g._allowed("https://graph.microsoft.com/v1.0/x")
    assert not g._allowed("http://graph.microsoft.com/v1.0/x")
    assert not g._allowed("https://graph.microsoft.com.evil.example/x")
    assert not g._allowed("https://contoso.sharepoint.com/x")  # only as a redirect target
    assert g._allowed("https://contoso.sharepoint.com/x", redirect=True)


@LO
def test_libreoffice_odf_roundtrip(tmp_path):
    lo = cn.LibreOfficeConnector()
    assert lo.status().available
    src = tmp_path / "note.docx"
    src.write_bytes(_docx_bytes("ODF round trip"))
    odt = lo.from_ooxml(src, tmp_path / "odf")
    assert odt.suffix == ".odt" and odt.stat().st_size > 0
    meta, back = lo.to_ooxml(odt, tmp_path / "back")
    import docx
    assert meta.format == "docx"
    assert "ODF round trip" in [p.text for p in docx.Document(str(back)).paragraphs]


def test_api_rejects_direct_import_template(st, tmp_path):
    from fastapi.testclient import TestClient

    from daedelus.api import create_app
    c = TestClient(create_app(tmp_path / "apiws"))
    pid = c.post("/api/projects", json={"name": "p"}).json()["id"]
    r = c.post(f"/api/projects/{pid}/artifacts", json={
        "name": "x", "adapter": "document", "template": "import",
        "params": {"path": "/etc/passwd"}})
    assert r.status_code == 400
    conns = {x["name"]: x for x in c.get("/api/connectors").json()}
    assert set(conns) == {"msgraph", "google", "libreoffice"}
    assert json.dumps(conns)  # serialisable status
