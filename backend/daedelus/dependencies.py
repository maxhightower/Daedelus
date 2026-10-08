"""Component-level artifact dependencies (V1.2).

A dependency says that one component of an artifact is *derived from* a component of another
artifact - e.g. a Word table shows the values of an Excel named range, a slide chart plots
them, a report figure renders a workbook chart, a slide's text summarises a report section.

This is distinct from:
- source bindings (a source *informs* an artifact: reference, inspiration, constraint...);
- workflow edges (execution order and data flow between workflow nodes).

Status is computed, never assumed: a dependency is *stale* when the source component's
authored state or calculated values differ from what was last synchronised. Syncing updates
only stale dependents, through the same checked edit path as manual edits (checkpoint, scope
and preservation checks, native validation, revision with ``origin="dependency"``); every
other component of the target keeps its exact state.
"""

from __future__ import annotations

import re
import tempfile
import threading
from pathlib import Path
from typing import Any, Literal

from pydantic import BaseModel, Field

from .adapters import get_adapter
from .artifacts import native_path
from .editing import EditRejected, apply_manual_edit
from .models import PlannedOperation, new_id, now_iso
from .store import ProjectStore

DOC_KIND = "dependencies"
ROUTES = {
    # (source kind, target kind) -> how the target is updated
    ("range", "table"): "values -> table rows",
    ("table", "table"): "values -> table rows",
    ("range", "chart"): "values -> chart categories/series",
    ("table", "chart"): "values -> chart categories/series",
    ("chart", "figure"): "rendered chart -> figure image",
    ("chart", "image"): "rendered chart -> picture",
    ("range", "paragraph"): "values -> text template",
    ("range", "text"): "values -> text template",
    ("range", "title"): "values -> text template",
    ("section", "text"): "section paragraphs -> text",
    ("section", "paragraph"): "section paragraphs -> text",
    ("paragraph", "text"): "paragraph -> text",
}
CREATE_OPS = {"table": "add_table", "chart": "add_chart", "figure": "add_figure",
              "image": "add_image", "paragraph": "add_paragraph", "text": "add_textbox"}


class End(BaseModel):
    artifact_id: str
    component_id: str


class ArtifactDependency(BaseModel):
    id: str = Field(default_factory=lambda: new_id("dep"))
    source: End
    target: End
    target_kind: str | None = None  # when the target does not exist yet (created on sync)
    options: dict[str, Any] = Field(default_factory=dict)
    # options: template (text), columns (chart series), number_format, max_paragraphs,
    #          max_chars, create {op params for creating a missing target}
    note: str = ""
    synced: dict[str, Any] | None = None
    created_at: str = Field(default_factory=now_iso)


Status = Literal["synced", "stale", "never_synced", "missing_source", "missing_target",
                 "needs_recalc", "unsupported"]


# ---------------------------------------------------------------------------- storage
def list_dependencies(store: ProjectStore) -> list[ArtifactDependency]:
    return store.list_docs(DOC_KIND, ArtifactDependency)


def save_dependency(store: ProjectStore, dep: ArtifactDependency) -> ArtifactDependency:
    store.put_doc(DOC_KIND, dep.id, dep)
    return dep


def delete_dependency(store: ProjectStore, dep_id: str) -> None:
    store.delete_doc(DOC_KIND, dep_id)


def add_dependency(store: ProjectStore, source: End, target: End, *, options: dict | None = None,
                   target_kind: str | None = None, note: str = "") -> ArtifactDependency:
    src_art = store.get_artifact(source.artifact_id)
    tgt_art = store.get_artifact(target.artifact_id)
    sc = src_art.component(source.component_id)
    if sc is None:
        raise ValueError(f"source component '{source.component_id}' not in {src_art.name}")
    tc = tgt_art.component(target.component_id)
    tk = tc.kind if tc else target_kind
    if tc is None and not (options or {}).get("create"):
        raise ValueError(f"target component '{target.component_id}' not in {tgt_art.name} "
                         "(give options.create to create it on first sync)")
    if (sc.kind, tk) not in ROUTES:
        raise ValueError(f"no dependency route from {sc.kind} to {tk}; supported: "
                         + ", ".join(f"{a}->{b}" for a, b in ROUTES))
    if source.artifact_id == target.artifact_id:
        raise ValueError("source and target must be different artifacts")
    for d in list_dependencies(store):
        if d.source == source and d.target == target:
            return d
    if _creates_cycle(store, source.artifact_id, target.artifact_id):
        raise ValueError("dependency would create a cycle between artifacts")
    return save_dependency(store, ArtifactDependency(source=source, target=target,
                                                     options=options or {}, target_kind=tk,
                                                     note=note))


def _creates_cycle(store: ProjectStore, src: str, tgt: str) -> bool:
    edges: dict[str, set[str]] = {}
    for d in list_dependencies(store):
        edges.setdefault(d.source.artifact_id, set()).add(d.target.artifact_id)
    edges.setdefault(src, set()).add(tgt)
    seen, stack = set(), [tgt]
    while stack:
        a = stack.pop()
        if a == src:
            return True
        if a in seen:
            continue
        seen.add(a)
        stack.extend(edges.get(a, ()))
    return False


# ---------------------------------------------------------------------------- status
def _inspect(store: ProjectStore, artifact_id: str, cache: dict[str, Any]):
    if artifact_id not in cache:
        art = store.get_artifact(artifact_id)
        ad = get_adapter(art.adapter)
        cache[artifact_id] = (art, ad, ad.inspect(native_path(store, art), art.entry))
    return cache[artifact_id]


def _signature(insp, art, cid: str) -> str:
    from .adapters.office.common import h

    sub = [cid] + art.descendants(cid)
    return h({c: (insp.states.get(c), insp.measurements.get(c, {}).get("values_hash"))
              for c in sub})


def status(store: ProjectStore, dep: ArtifactDependency, cache: dict | None = None
           ) -> dict[str, Any]:
    cache = {} if cache is None else cache
    src_art, _, src_insp = _inspect(store, dep.source.artifact_id, cache)
    tgt_art, _, tgt_insp = _inspect(store, dep.target.artifact_id, cache)
    sc = next((c for c in src_insp.components if c.id == dep.source.component_id), None)
    tc = next((c for c in tgt_insp.components if c.id == dep.target.component_id), None)
    out: dict[str, Any] = {"id": dep.id, "route": ROUTES.get(
        (sc.kind if sc else None, tc.kind if tc else dep.target_kind)),
        "source": {**dep.source.model_dump(), "artifact_name": src_art.name,
                   "component_name": sc.name if sc else None, "kind": sc.kind if sc else None},
        "target": {**dep.target.model_dump(), "artifact_name": tgt_art.name,
                   "component_name": tc.name if tc else None,
                   "kind": tc.kind if tc else dep.target_kind}}
    if sc is None:
        out.update(status="missing_source", reason="source component no longer exists")
        return out
    if tc is None and not dep.options.get("create"):
        out.update(status="missing_target", reason="target component no longer exists")
        return out
    m = src_insp.measurements.get(sc.id, {})
    if sc.kind in ("range", "table") and m.get("calculated") is False:
        out.update(status="needs_recalc", reason="source formulas are not calculated "
                   "(LibreOffice unavailable); values cannot be propagated")
        return out
    sig = _signature(src_insp, src_art, sc.id)
    out["source_signature"] = sig
    if tc is None:
        out.update(status="never_synced", reason="target will be created on first sync")
    elif dep.synced is None:
        out.update(status="never_synced", reason="not synchronised yet")
    elif dep.synced.get("source_signature") != sig:
        out.update(status="stale", reason="source changed since the last sync "
                   f"(synced from revision {dep.synced.get('source_revision_id')})")
    else:
        out.update(status="synced", reason=f"in sync with source revision "
                   f"{dep.synced.get('source_revision_id')}")
    return out


_STATUS_CACHE: dict[str, tuple[str, list[dict[str, Any]]]] = {}


def cached_status(store: ProjectStore) -> list[dict[str, Any]]:
    """all_status, recomputed only when a dependency or an involved artifact head changed."""
    deps = list_dependencies(store)
    if not deps:
        return []
    arts = sorted({d.source.artifact_id for d in deps} | {d.target.artifact_id for d in deps})
    heads = []
    for a in arts:
        try:
            heads.append(store.get_artifact(a).head_revision_id or "")
        except LookupError:
            heads.append("missing")
    key = _h_json([d.model_dump() for d in deps] + heads)
    root = str(store.root)
    hit = _STATUS_CACHE.get(root)
    if hit and hit[0] == key:
        return hit[1]
    val = all_status(store)
    _STATUS_CACHE[root] = (key, val)
    return val


def _h_json(v: Any) -> str:
    import hashlib
    import json
    return hashlib.sha256(json.dumps(v, sort_keys=True, default=str).encode()).hexdigest()


def all_status(store: ProjectStore) -> list[dict[str, Any]]:
    cache: dict[str, Any] = {}
    return [status(store, d, cache) for d in list_dependencies(store)]


# ---------------------------------------------------------------------------- materialise
CELL_RE = re.compile(r"\{([A-Z]{1,3})([0-9]{1,5})(?::([^}]+))?\}")


def _grid_cell(grid: list[list[Any]], col: str, row: int, origin: tuple[int, int]) -> Any:
    from openpyxl.utils import column_index_from_string

    c = column_index_from_string(col) - origin[0]
    r = row - origin[1]
    try:
        return grid[r][c]
    except IndexError:
        raise ValueError(f"template cell {col}{row} lies outside the source range") from None


def _fmt(v: Any, spec: str | None) -> str:
    if v is None:
        return "?"
    if spec:
        return format(v, spec)
    if isinstance(v, float):
        return f"{v:,.2f}"
    return str(v)


def _source_values(store, dep, src_art, src_insp) -> tuple[list[list[Any]], tuple[int, int]]:
    from openpyxl.utils.cell import range_boundaries

    props = src_insp.properties.get(dep.source.component_id, {})
    vals = props.get("values")
    if vals is None:
        raise ValueError("source has no values")
    c1, r1, _, _ = range_boundaries(props["ref"].replace("$", ""))
    return vals, (c1, r1)


def _ops_for(store: ProjectStore, dep: ArtifactDependency, cache, files: dict[str, str],
             tmp: Path) -> list[PlannedOperation]:
    src_art, src_ad, src_insp = _inspect(store, dep.source.artifact_id, cache)
    tgt_art, _, tgt_insp = _inspect(store, dep.target.artifact_id, cache)
    sk = next(c.kind for c in src_insp.components if c.id == dep.source.component_id)
    tc = next((c for c in tgt_insp.components if c.id == dep.target.component_id), None)
    tk = tc.kind if tc else dep.target_kind
    opt = dep.options
    tid = dep.target.component_id
    create = None if tc is not None else dict(opt.get("create") or {})
    if create is not None:
        create["id"] = tid

    def op(name, params):
        if create is not None:
            return PlannedOperation(op=CREATE_OPS[tk], params={**create, **params},
                                    rationale=f"create dependent {tk} from {dep.source.component_id}")
        return PlannedOperation(op=name, component_id=tid, params=params,
                                rationale=f"dependency {dep.id}: {ROUTES[(sk, tk)]}")

    if sk in ("range", "table") and tk == "table":
        vals, _ = _source_values(store, dep, src_art, src_insp)
        rows = [[_fmt(v, opt.get("number_format")) if isinstance(v, float) else
                 ("" if v is None else v) for v in r] for r in vals]
        return [op("update_table", {"rows": rows})]
    if sk in ("range", "table") and tk == "chart":
        vals, _ = _source_values(store, dep, src_art, src_insp)
        header, body = vals[0], vals[1:]
        cols = opt.get("columns") or list(range(1, len(header)))
        cats = [str(r[0]) for r in body]
        series = {str(header[c]): [None if r[c] is None else float(r[c]) for r in body]
                  for c in cols}
        return [op("update_chart_data", {"categories": cats, "series": series})]
    if sk == "chart" and tk in ("figure", "image"):
        out = tmp / f"{dep.id}.png"
        src_ad.render_chart(native_path(store, src_art), src_art.entry,
                            dep.source.component_id, out)
        key = f"dep_{dep.id}"
        files[key] = str(out)
        if create is not None:
            return [PlannedOperation(op=CREATE_OPS[tk], params={**create, "image": key},
                                     rationale="create figure from rendered chart")]
        name = "set_figure_image" if tk == "figure" else "set_image"
        return [PlannedOperation(op=name, component_id=tid, params={"image": key},
                                 rationale=f"dependency {dep.id}: re-rendered chart")]
    if sk in ("range", "table") and tk in ("paragraph", "text", "title"):
        vals, origin = _source_values(store, dep, src_art, src_insp)
        tpl = opt.get("template")
        if not tpl:
            raise ValueError("a values -> text dependency needs options.template")
        text = CELL_RE.sub(lambda m: _fmt(_grid_cell(vals, m.group(1), int(m.group(2)), origin),
                                          m.group(3)), tpl)
        return [op("set_text", {"text": text})]
    if sk in ("section", "paragraph") and tk in ("text", "paragraph"):
        st = src_insp.properties
        sub = [dep.source.component_id] + src_art.descendants(dep.source.component_id) \
            if sk == "section" else [dep.source.component_id]
        paras = [st[c]["text"] for c in sub if c in st and st[c].get("text")
                 and next((x.kind for x in src_insp.components if x.id == c), "") == "paragraph"]
        n = int(opt.get("max_paragraphs", 4))
        mx = int(opt.get("max_chars", 180))
        paras = [p if len(p) <= mx else p[:mx].rsplit(" ", 1)[0] + "…" for p in paras[:n]]
        if tk == "paragraph":
            return [op("set_text", {"text": " ".join(paras)})]
        return [op("set_text", {"paragraphs": paras})]
    raise ValueError(f"no route {sk} -> {tk}")


def sync(store: ProjectStore, *, dependency_ids: list[str] | None = None,
         target_artifact_ids: list[str] | None = None, force: bool = False,
         lock: threading.RLock | None = None) -> dict[str, Any]:
    """Update stale dependents. One checked revision per target artifact."""
    deps = [d for d in list_dependencies(store)
            if (dependency_ids is None or d.id in dependency_ids)
            and (target_artifact_ids is None or d.target.artifact_id in target_artifact_ids)]
    cache: dict[str, Any] = {}
    report = {"checked": len(deps), "updated": [], "skipped": [], "failed": [], "revisions": []}
    todo: dict[str, list[tuple[ArtifactDependency, dict]]] = {}
    for d in deps:
        st = status(store, d, cache)
        if st["status"] in ("stale", "never_synced") or (force and st["status"] == "synced"):
            todo.setdefault(d.target.artifact_id, []).append((d, st))
        else:
            report["skipped"].append({"id": d.id, "status": st["status"],
                                      "reason": st.get("reason")})
    with tempfile.TemporaryDirectory(prefix="dd_dep_") as td:
        for tgt_id, items in todo.items():
            files: dict[str, str] = {}
            ops: list[PlannedOperation] = []
            ok_items = []
            for d, st in items:
                try:
                    ops += _ops_for(store, d, cache, files, Path(td))
                    ok_items.append((d, st))
                except Exception as exc:  # report precisely; other dependents still sync
                    report["failed"].append({"id": d.id, "error": f"{type(exc).__name__}: {exc}"})
            if not ops:
                continue
            art = store.get_artifact(tgt_id)
            try:
                rev = apply_manual_edit(
                    store, tgt_id, ops, base_revision_id=art.head_revision_id, lock=lock,
                    files=files, origin="dependency",
                    message="Update dependents: " + ", ".join(
                        f"{d.target.component_id} <- {store.get_artifact(d.source.artifact_id).name}"
                        f"/{d.source.component_id}" for d, _ in ok_items))
            except EditRejected as exc:
                report["failed"] += [{"id": d.id, "error": str(exc)} for d, _ in ok_items]
                continue
            report["revisions"].append({"artifact_id": tgt_id, "revision_id": rev.id,
                                        "revision_number": rev.number,
                                        "changed": rev.changed_components})
            cache.pop(tgt_id, None)
            for d, st in ok_items:
                src = store.get_artifact(d.source.artifact_id)
                d.synced = {"source_signature": st["source_signature"],
                            "source_revision_id": src.head_revision_id,
                            "target_revision_id": rev.id, "at": now_iso()}
                save_dependency(store, d)
                report["updated"].append({"id": d.id, "target": d.target.model_dump(),
                                          "route": st["route"]})
    return report


def dependents_of(store: ProjectStore, artifact_id: str) -> list[ArtifactDependency]:
    return [d for d in list_dependencies(store) if d.source.artifact_id == artifact_id]
