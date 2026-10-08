"""V1.2 demonstration: the Research Analysis and Presentation Pipeline.

Inputs (``demo_assets/v12``): the NOAA GML Mauna Loa CO2 annual means (real data, file kept
verbatim), research notes and a visual guide written for the demonstration.

1. One workflow builds three native, editable files from the bound sources:
   an analysis workbook (cleaned data, formulas, named ranges, an Excel table, two native
   charts), a report (sections quoted from the notes, a results table, a figure rendered from
   the workbook chart, references) and a slide deck (theme from the guide, a native chart and
   findings text) - then a ``dependencies`` node links components across the three files.
2. The workbook's formula results are checked against values computed independently from the
   CSV in Python (LibreOffice recalculates; openpyxl never evaluates formulas).
3. Native validation of all three files (reopen with the libraries, open in LibreOffice, ids).
4. A **demonstration edit** (not NOAA data) changes the 2020s rows of the data range; the
   dependency status shows exactly which dependents went stale; "Update dependents" changes
   only those components - everything else is compared state by state.

Verification type: integration (real openpyxl / python-docx / python-pptx / LibreOffice,
deterministic local provider). No AI model is involved: the authoring recipes are fixed
templates parameterised by the source contents and labelled as such.

Run: ``daedelus demo-v12 --out evidence/v1_2``
"""

from __future__ import annotations

import csv
import io
import json
import shutil
import time
from pathlib import Path
from typing import Any

from . import connectors as conn_mod
from . import dependencies as dp
from . import ingest
from .adapters import get_adapter
from .artifacts import create_artifact, native_path
from .editing import apply_manual_edit
from .engine import Engine
from .models import (PlannedOperation, SourceBinding, TargetSelector, Workflow, WorkflowEdge,
                     WorkflowNode)
from .scenarios_v11 import Demo, _dump
from .store import ProjectStore, Workspace

ASSETS = Path(__file__).with_name("demo_assets") / "v12"
DEMO_SHIFT = 10.0  # ppm added to the 2020s rows by the demonstration edit (synthetic change)


def _expected_decades(csv_path: Path, shift_from: int | None = None) -> dict[str, float]:
    """Decade means computed directly from the CSV, independent of Daedelus and LibreOffice."""
    text = "\n".join(ln for ln in csv_path.read_text().splitlines() if not ln.startswith("#"))
    rows = list(csv.DictReader(io.StringIO(text)))
    acc: dict[int, list[float]] = {}
    for r in rows:
        y, v = int(r["year"]), float(r["mean"])
        if shift_from is not None and y >= shift_from:
            v += DEMO_SHIFT
        acc.setdefault(y // 10 * 10, []).append(v)
    return {f"{d}s": sum(v) / len(v) for d, v in sorted(acc.items())}


def _states(st: ProjectStore, aid: str) -> dict[str, str]:
    a = st.get_artifact(aid)
    return dict(get_adapter(a.adapter).inspect(native_path(st, a), a.entry).states)


def _copy_native(st: ProjectStore, aid: str, dst: Path, name: str) -> Path:
    a = st.get_artifact(aid)
    dst.mkdir(parents=True, exist_ok=True)
    out = dst / name
    shutil.copy2(native_path(st, a) / a.entry, out)
    return out


def _copy_previews(st: ProjectStore, aid: str, dst: Path, prefix: str) -> list[str]:
    a = st.get_artifact(aid)
    rev = st.get_revision(a.head_revision_id)
    dst.mkdir(parents=True, exist_ok=True)
    out = []
    for k, rel in sorted(rev.previews.items()):
        p = st.abs(rel)
        if p.is_file() and p.suffix in (".png", ".pdf"):
            shutil.copy2(p, dst / f"{prefix}_{k}{p.suffix}")
            out.append(f"{prefix}_{k}{p.suffix}")
    return out


def build_pipeline(st: ProjectStore) -> dict[str, Any]:
    """Sources, three empty artifacts, bindings and the pipeline workflow (also used by e2e)."""
    data = ingest.register_file(st, ASSETS / "co2_annmean_mlo.csv",
                                name="Mauna Loa CO2 annual means (NOAA GML)")
    notes = ingest.register_file(st, ASSETS / "research_notes.md", name="Research notes")
    guide = ingest.register_file(st, ASSETS / "visual_guide.md", name="Visual guide")
    wb, _ = create_artifact(st, name="CO2 analysis", adapter="spreadsheet", template="blank",
                            params={"name": "CO2 analysis"})
    doc, _ = create_artifact(st, name="CO2 report", adapter="document", template="blank",
                             params={"title": "Atmospheric CO2 at Mauna Loa", "name": "CO2 report"})
    deck, _ = create_artifact(st, name="CO2 briefing", adapter="presentation", template="blank",
                              params={"name": "CO2 briefing"})

    def bind(s, a, role, aspects):
        st.save_binding(SourceBinding(source_id=s.id, role=role, aspects=aspects,
                                      target=TargetSelector(scope="artifact", artifact_id=a.id)))
    bind(data, wb, "reference", ["data"])
    bind(notes, doc, "reference", ["content", "structure"])
    bind(data, doc, "context", ["data"])
    bind(guide, deck, "guideline", ["style", "color"])
    bind(notes, deck, "context", ["content"])
    links = [
        {"source": {"artifact_id": wb.id, "component_id": "summary_values"},
         "target": {"artifact_id": doc.id, "component_id": "results_table"},
         "target_kind": "table", "options": {"create": {"parent": "results",
                                                        "after": "results_text"}},
         "note": "decade summary -> report table"},
        {"source": {"artifact_id": wb.id, "component_id": "decade_chart"},
         "target": {"artifact_id": doc.id, "component_id": "fig_decades"},
         "target_kind": "figure", "options": {"create": {
             "parent": "results", "after": "results_table", "width_cm": 14,
             "caption": "Figure 1. Mean CO2 by decade (rendered from the workbook chart)"}},
         "note": "workbook chart -> report figure"},
        {"source": {"artifact_id": wb.id, "component_id": "summary_values"},
         "target": {"artifact_id": doc.id, "component_id": "results_text"},
         "options": {"template": "Decade means rose from {B2:.1f} ppm in the {A2} to {B9:.1f} "
                                 "ppm in the {A9} (computed in the linked workbook)."},
         "note": "named range -> results sentence"},
        {"source": {"artifact_id": wb.id, "component_id": "decade_means"},
         "target": {"artifact_id": deck.id, "component_id": "deck_chart"},
         "target_kind": "chart", "options": {"create": {
             "slide": "data_slide", "box": [0.08, 0.22, 0.84, 0.7], "type": "column",
             "title": "Mean CO2 by decade (ppm)"}},
         "note": "named range -> native slide chart"},
        {"source": {"artifact_id": doc.id, "component_id": "discussion"},
         "target": {"artifact_id": deck.id, "component_id": "findings.body"},
         "options": {"max_paragraphs": 2}, "note": "report section -> slide text"},
    ]

    def agent(nid, aid, label, instr):
        return [WorkflowNode(id=f"art_{nid}", type="artifact", label=label,
                             config={"artifact_id": aid}),
                WorkflowNode(id=nid, type="agent", label=f"{label} agent", config={
                    "instructions": instr, "provider": "heuristic", "fan_out": False})]
    nodes = [WorkflowNode(id="src", type="sources", label="Sources", config={})]
    nodes += agent("build_wb", wb.id, "Workbook",
                   "Build the analysis workbook from the dataset.\nValue label: CO2 (ppm)")
    nodes += agent("build_doc", doc.id, "Report", "Write the report from the research notes.")
    nodes += agent("build_deck", deck.id, "Deck",
                   "Title: Atmospheric CO2 at Mauna Loa\nBuild the slide deck.")
    nodes.append(WorkflowNode(id="deps", type="dependencies", label="Link components",
                              config={"links": links}))
    nodes.append(WorkflowNode(id="validate", type="validate", label="Validate", config={
        "checks": ["file_reopens", "component_preservation"], "fail_on_error": True}))
    edges = []
    for n in ("build_wb", "build_doc", "build_deck"):
        edges += [WorkflowEdge(id=f"a_{n}", source=f"art_{n}", source_port="artifact", target=n,
                               target_port="artifact"),
                  WorkflowEdge(id=f"s_{n}", source="src", source_port="sources", target=n,
                               target_port="sources"),
                  WorkflowEdge(id=f"d_{n}", source=n, source_port="revision", target="deps",
                               target_port="after"),
                  WorkflowEdge(id=f"v_{n}", source=n, source_port="revision", target="validate",
                               target_port="revision")]
    wf = Workflow(name="Research analysis and presentation", nodes=nodes, edges=edges)
    st.save_workflow(wf)
    return {"wf": wf, "wb": wb, "doc": doc, "deck": deck, "sources": [data, notes, guide]}


def run(out: Path, workspace: Path | None = None) -> dict[str, Any]:
    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)
    ws = Workspace(workspace or out / "workspace")
    _, st = ws.create_project("V1.2 research pipeline", "Research analysis and presentation")
    t0 = time.perf_counter()
    d = Demo("P", "Research analysis and presentation pipeline", "integration")
    ctx = build_pipeline(st)
    wb, doc, deck = ctx["wb"], ctx["doc"], ctx["deck"]
    ex = Engine(st).execute(ctx["wf"].id)
    d.metrics["build_seconds"] = round(time.perf_counter() - t0, 1)
    d.check("pipeline workflow succeeded", ex.status.value == "succeeded", ex.error)
    for n in ("build_wb", "build_doc", "build_deck", "deps", "validate"):
        nr = ex.run(n)
        d.check(f"node {n} succeeded", nr.status.value == "succeeded", nr.error)
    _dump(out / "plans.json", {n: [u.plan for u in ex.run(n).units]
                               for n in ("build_wb", "build_doc", "build_deck")})
    d.notes += [n for u in ex.run("build_wb").units for n in (u.plan or {}).get("notes", [])]

    # -------------------------------------------------------------- workbook checks
    w = st.get_artifact(wb.id)
    ids = {c.id: c.kind for c in w.components}
    d.check("workbook hierarchy: sheets, table, named ranges, charts",
            {"summary": "sheet", "data_table": "table", "dataset": "range",
             "summary_values": "range", "decade_means": "range", "trend_chart": "chart",
             "decade_chart": "chart"}.items() <= ids.items(), ids)
    import openpyxl
    wbx = openpyxl.load_workbook(native_path(st, w) / w.entry)
    f = wbx["Summary"]["B2"].value
    d.check("summary values are live formulas in the native file",
            isinstance(f, str) and f.startswith("=AVERAGEIFS("), f)
    sa = get_adapter("spreadsheet")
    vals = sa.values(native_path(st, w), w.entry, "Summary!A1:B9")
    expected = _expected_decades(ASSETS / "co2_annmean_mlo.csv")
    got = {r[0]: r[1] for r in vals[1:]}
    diffs = {k: (got.get(k), round(v, 4)) for k, v in expected.items()
             if got.get(k) is None or abs(got[k] - v) > 1e-6}
    d.check("formula results (LibreOffice recalculation) match decade means computed "
            "independently from the CSV", not diffs, diffs)
    d.metrics["decade_means"] = {k: round(v, 3) for k, v in expected.items()}
    with __import__("zipfile").ZipFile(native_path(st, w) / w.entry) as z:
        n_charts = len([n for n in z.namelist() if n.startswith("xl/charts/chart")])
    d.check("two native Excel charts in the workbook package", n_charts == 2, n_charts)

    # -------------------------------------------------------------- report checks
    r = st.get_artifact(doc.id)
    rk = {c.id: c.kind for c in r.components}
    d.check("report hierarchy: title, sections, paragraphs, results table, figure, references",
            {"summary": "section", "data_method": "section", "results": "section",
             "discussion": "section", "references": "section", "results_table": "table",
             "fig_decades": "figure", "results_text": "paragraph"}.items() <= rk.items(), rk)
    da = get_adapter("document")
    rp = da.inspect(native_path(st, r), r.entry).properties
    tbl = rp["results_table"]["rows"]
    d.check("report table carries the workbook's calculated values",
            tbl[1][0] == "1950s" and abs(float(tbl[1][1]) - expected["1950s"]) < 0.01, tbl[:3])
    d.check("results sentence filled from the named range",
            f"{expected['2020s']:.1f}" in rp["results_text"]["text"], rp["results_text"]["text"])
    notes_text = (ASSETS / "research_notes.md").read_text()
    d.check("report body paragraphs are quoted from the research notes (no generated prose)",
            rp["summary_text"]["text"] in " ".join(notes_text.split()), rp["summary_text"]["text"])

    # -------------------------------------------------------------- deck checks
    k = st.get_artifact(deck.id)
    from pptx import Presentation
    prs = Presentation(str(native_path(st, k) / k.entry))
    kinds = [[("chart" if s.has_chart else "picture" if s.shape_type == 13 else
               "text" if s.has_text_frame else "shape") for s in sl.shapes] for sl in prs.slides]
    d.check("4 slides built from editable objects (text frames, native chart, shapes; "
            "no slide is a picture)", len(kinds) == 4 and all(
                "picture" not in x for x in kinds) and any("chart" in x for x in kinds), kinds)
    chart = next(s.chart for sl in prs.slides for s in sl.shapes if s.has_chart)
    cvals = list(chart.plots[0].series[0].values)
    d.check("slide chart data comes from the workbook (data-driven, editable)",
            abs(cvals[-1] - expected["2020s"]) < 0.01, cvals)
    title_run = prs.slides[0].shapes.title.text_frame.paragraphs[0].runs[0]
    d.check("theme from the visual guide applied (title colour and font)",
            str(title_run.font.color.rgb).lower() == "0b3d5c" and title_run.font.name == "Georgia",
            (str(title_run.font.color.rgb), title_run.font.name))

    # -------------------------------------------------------------- dependencies + validation
    st0 = dp.all_status(st)
    d.check("5 component dependencies, all synced after the run",
            len(st0) == 5 and all(x["status"] == "synced" for x in st0),
            [(x["target"]["component_id"], x["status"]) for x in st0])
    vrep = {}
    for a in (w, r, k):
        a = st.get_artifact(a.id)
        rep = get_adapter(a.adapter).validate(native_path(st, a), a.entry, ["file_reopens"], {})
        vrep[a.name] = [c.model_dump() for c in rep.checks]
        d.check(f"native validation of {a.name} ({a.entry})", rep.passed,
                [c.name for c in rep.checks if not c.passed])
    _dump(out / "validation.json", vrep)

    ev = out / "before"
    files = {"workbook": _copy_native(st, wb.id, ev, "CO2_analysis.xlsx"),
             "report": _copy_native(st, doc.id, ev, "CO2_report.docx"),
             "deck": _copy_native(st, deck.id, ev, "CO2_briefing.pptx")}
    for key, aid in (("workbook", wb.id), ("report", doc.id), ("deck", deck.id)):
        _copy_previews(st, aid, ev / "previews", key)
    d.files.update({k2: str(v) for k2, v in files.items()})
    _dump(out / "dependencies_before.json", st0)

    # -------------------------------------------------------------- selective update
    before = {a: _states(st, a) for a in (doc.id, deck.id)}
    w = st.get_artifact(wb.id)
    data_sheet = next(c for c in w.components if c.kind == "sheet" and c.name == "Data")
    rows = sa.values(native_path(st, w), w.entry, "Data!A1:B200")
    first = next(i for i, row in enumerate(rows) if isinstance(row[0], int) and row[0] >= 2020)
    last = max(i for i, row in enumerate(rows) if isinstance(row[0], (int, float)))
    new_vals = [[round(rows[i][1] + DEMO_SHIFT, 3)] for i in range(first, last + 1)]
    t1 = time.perf_counter()
    rev = apply_manual_edit(st, wb.id, [PlannedOperation(
        op="write_range", component_id=data_sheet.id,
        params={"start": f"B{first + 1}", "values": new_vals})],
        message=f"Demonstration edit: 2020s values +{DEMO_SHIFT} ppm (synthetic, not NOAA data)")
    d.check("data range edit is a checked revision of the workbook", rev.number >= 3, rev.number)
    stale = {x["target"]["component_id"]: x["status"] for x in dp.all_status(st)}
    d.check("exactly the dependents of the changed range went stale",
            stale == {"results_table": "stale", "fig_decades": "stale", "results_text": "stale",
                      "deck_chart": "stale", "findings.body": "synced"}, stale)
    _dump(out / "dependencies_stale.json", dp.all_status(st))
    rep = dp.sync(st, lock=None)
    d.metrics["update_seconds"] = round(time.perf_counter() - t1, 1)
    changed = {st.get_artifact(x["artifact_id"]).name: sorted(x["changed"])
               for x in rep["revisions"]}
    d.check("update created one revision per dependent artifact",
            len(rep["revisions"]) == 2 and not rep["failed"], rep)
    after = {a: _states(st, a) for a in (doc.id, deck.id)}
    moved = {a: sorted(c for c in after[a] if after[a].get(c) != before[a].get(c))
             for a in after}
    d.check("report: only the table, figure and results sentence changed",
            set(moved[doc.id]) - {"document"} == {"results_table", "fig_decades",
                                                 "results_text"}, moved[doc.id])
    d.check("deck: only the data chart changed (findings text untouched)",
            set(moved[deck.id]) - {"presentation", "data_slide"} == {"deck_chart"},
            moved[deck.id])
    exp2 = _expected_decades(ASSETS / "co2_annmean_mlo.csv", shift_from=2020)
    r = st.get_artifact(doc.id)
    rp = da.inspect(native_path(st, r), r.entry).properties
    d.check("updated sentence shows the recalculated 2020s mean",
            f"{exp2['2020s']:.1f}" in rp["results_text"]["text"], rp["results_text"]["text"])
    d.check("all dependencies synced again",
            all(x["status"] == "synced" for x in dp.all_status(st)))
    _dump(out / "update_report.json", {"changed": changed, "moved_states": {
        st.get_artifact(a).name: v for a, v in moved.items()}, "sync": rep})
    ev2 = out / "after"
    _copy_native(st, wb.id, ev2, "CO2_analysis.xlsx")
    _copy_native(st, doc.id, ev2, "CO2_report.docx")
    _copy_native(st, deck.id, ev2, "CO2_briefing.pptx")
    for key, aid in (("workbook", wb.id), ("report", doc.id), ("deck", deck.id)):
        _copy_previews(st, aid, ev2 / "previews", key)
    d.notes.append(f"The data change is a demonstration edit (+{DEMO_SHIFT} ppm on 2020-2025), "
                   "not NOAA data; it exists only to show selective dependency updates.")

    report = {"seconds": round(time.perf_counter() - t0, 1), "demos": [d.dump()],
              "connectors": [c.status().model_dump() for c in conn_mod.registry().values()],
              "workspace": str(st.root)}
    _dump(out / "demo_v12_report.json", report)
    lines = ["# V1.2 demonstration", "", f"Duration {report['seconds']} s.", "",
             f"## {d.title}", "", f"Status: **{d.status.upper()}** - verification type: "
             f"{d.verification}", ""]
    lines += [f"- [{'x' if c['ok'] else ' '}] {c['name']}" + (f" - {c['detail']}" if not c["ok"]
                                                              else "") for c in d.checks]
    lines += ["", "Notes:"] + [f"- {n}" for n in d.notes] + ["", "Connectors:"]
    lines += [f"- {c['title']}: {'available' if c['available'] else c['reason']} "
              f"({c['verification']})" for c in report["connectors"]]
    (out / "demo_v12_report.md").write_text("\n".join(lines) + "\n")
    ws.close()
    return report
