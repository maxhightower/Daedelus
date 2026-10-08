"""Word document adapter (.docx via python-docx).

Components: document -> sections (a heading and the content up to the next heading of the
same or higher level) -> paragraphs / tables / figures. Identity uses native Word bookmarks
named ``dd_<component id>`` (they survive editing in Word and LibreOffice). Elements without a
bookmark (e.g. typed by a person in Word) are listed as *unregistered* and receive a bookmark
the next time Daedelus saves the document. Missing or duplicated bookmarks are reported and
operations on them are refused.

python-docx keeps XML it does not understand, so unrelated content survives targeted edits;
the engine's component-preservation check verifies that on every revision.
"""

from __future__ import annotations

import copy
import difflib
import hashlib
import json
import shutil
from pathlib import Path
from typing import Any

from ...models import Component, PlannedOperation, ValidationReport
from ..base import Adapter, AdapterError, AdapterInfo, ApplyResult, InspectResult, OperationSpec
from . import common as oc

ENTRY = "document.docx"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
ID = {"type": "string", "pattern": "[A-Za-z0-9_.-]+"}
PLACE = {"after": {"type": "string"}, "parent": {"type": "string"}}


def _docx():
    try:
        import docx  # noqa: F401
    except ImportError as exc:  # pragma: no cover
        raise AdapterError("python-docx is not installed") from exc
    import docx

    return docx


def _q(tag: str) -> str:
    p, t = tag.split(":")
    return "{%s}%s" % ({"w": W_NS, "a": A_NS, "r": R_NS}[p], t)


def _heading_level(p_el) -> int | None:
    st = p_el.find(f"{_q('w:pPr')}/{_q('w:pStyle')}")
    if st is None:
        return None
    val = st.get(_q("w:val")) or ""
    if val.lower().startswith("heading"):
        try:
            return int(val[7:])
        except ValueError:
            return 1
    if val.lower() == "title":
        return 0
    return None


def _bookmarks(el) -> list[str]:
    return [b.get(_q("w:name")) for b in el.iter(_q("w:bookmarkStart"))
            if (b.get(_q("w:name")) or "").startswith("dd_")]


def _text(el) -> str:
    return "".join(t.text or "" for t in el.iter(_q("w:t")))


def _has_picture(el) -> bool:
    return any(True for _ in el.iter(_q("a:blip")))


class DocumentAdapter(Adapter):
    name = "document"
    version = "1"

    def check_environment(self) -> tuple[bool, str]:
        try:
            _docx()
        except AdapterError as exc:
            return False, str(exc)
        return True, f"python-docx; LibreOffice {'found' if oc.soffice() else 'not found (no preview)'}"

    def info(self) -> AdapterInfo:
        ok, detail = self.check_environment()
        text_kinds = ["paragraph", "section"]
        return AdapterInfo(
            name=self.name, version=self.version,
            description="Word documents (.docx) edited with python-docx; bookmarks give "
                        "sections, paragraphs, tables and figures stable identities.",
            artifact_types=["document"], native_formats=[".docx"],
            preview_formats=["png pages (LibreOffice)", "structure (JSON)"],
            export_formats=["docx", "pdf", "md"],
            scopes=["document", "section", "paragraph", "table", "figure"],
            measurements=["words", "characters", "paragraphs", "rows", "columns", "figures"],
            validation=["file_reopens", "libreoffice_opens", "component_identity",
                        "component_preservation"],
            environment={"requires": "python-docx (LibreOffice optional for previews)"},
            available=ok, unavailable_reason=None if ok else detail, templates=["blank"],
            operations=[
                OperationSpec(name="add_heading", family="create", aspects=["structure", "content"],
                              target_kinds=["new"], description="Add a heading (starts a section) "
                              "at the end, after a component, or at the end of a parent section.",
                              params_schema={"type": "object", "required": ["id", "text"],
                                             "additionalProperties": False, "properties": {
                                                 "id": ID, "text": {"type": "string"},
                                                 "level": {"type": "integer", "minimum": 0, "maximum": 6},
                                                 **PLACE}}),
                OperationSpec(name="add_paragraph", family="create", aspects=["content", "text"],
                              target_kinds=["new"], description="Add a paragraph.",
                              params_schema={"type": "object", "required": ["id", "text"],
                                             "additionalProperties": False, "properties": {
                                                 "id": ID, "text": {"type": "string"},
                                                 "style": {"type": "string"}, **PLACE}}),
                OperationSpec(name="set_text", family="content", aspects=["content", "text"],
                              target_kinds=text_kinds, description="Replace the text of a paragraph "
                              "or section heading (formatting of the first run is kept).",
                              params_schema={"type": "object", "required": ["text"],
                                             "additionalProperties": False, "properties": {
                                                 "text": {"type": "string"},
                                                 "style": {"type": "string"}}}),
                OperationSpec(name="replace_section", family="content", aspects=["content", "text"],
                              target_kinds=["section"],
                              description="Replace the body paragraphs of a section (tables, "
                              "figures and sub-sections are kept).",
                              params_schema={"type": "object", "required": ["paragraphs"],
                                             "additionalProperties": False, "properties": {
                                                 "paragraphs": {"type": "array", "items": {
                                                     "type": "string"}, "maxItems": 200},
                                                 "style": {"type": "string"}}}),
                OperationSpec(name="add_table", family="create", aspects=["data", "content"],
                              target_kinds=["new"], description="Insert a table.",
                              params_schema={"type": "object", "required": ["id", "rows"],
                                             "additionalProperties": False, "properties": {
                                                 "id": ID, "rows": {"type": "array", "items": {
                                                     "type": "array"}},
                                                 "header": {"type": "boolean"},
                                                 "style": {"type": "string"},
                                                 "number_format": {"type": "string"}, **PLACE}}),
                OperationSpec(name="update_table", family="data", aspects=["data", "content"],
                              target_kinds=["table"], description="Replace a table's cell values "
                              "(rows are added/removed as needed; table style is kept).",
                              params_schema={"type": "object", "required": ["rows"],
                                             "additionalProperties": False, "properties": {
                                                 "rows": {"type": "array", "items": {
                                                     "type": "array"}},
                                                 "number_format": {"type": "string"}}}),
                OperationSpec(name="add_figure", family="create", aspects=["content", "figure"],
                              target_kinds=["new"], description="Insert an image with an optional "
                              "caption. 'image' is a source id or a file key provided by the engine.",
                              params_schema={"type": "object", "required": ["id", "image"],
                                             "additionalProperties": False, "properties": {
                                                 "id": ID, "image": {"type": "string"},
                                                 "width_cm": {"type": "number", "minimum": 1, "maximum": 30},
                                                 "caption": {"type": "string"}, **PLACE}}),
                OperationSpec(name="set_figure_image", family="figure", aspects=["figure"],
                              target_kinds=["figure"], description="Replace a figure's image "
                              "(keeps its position, caption and bookmark).",
                              params_schema={"type": "object", "required": ["image"],
                                             "additionalProperties": False, "properties": {
                                                 "image": {"type": "string"},
                                                 "width_cm": {"type": "number", "minimum": 1, "maximum": 30}}}),
                OperationSpec(name="apply_style", family="style", aspects=["style"],
                              target_kinds=text_kinds, description="Apply a paragraph style.",
                              params_schema={"type": "object", "required": ["style"],
                                             "additionalProperties": False,
                                             "properties": {"style": {"type": "string"}}}),
                OperationSpec(name="set_page_layout", family="layout", aspects=["layout"],
                              target_kinds=["document"], description="Orientation and margins.",
                              params_schema={"type": "object", "additionalProperties": False,
                                             "properties": {
                                                 "orientation": {"type": "string",
                                                                 "enum": ["portrait", "landscape"]},
                                                 "margins_cm": {"type": "number", "minimum": 0.5,
                                                                "maximum": 6}}}),
                OperationSpec(name="delete_component", family="structure", aspects=["structure"],
                              target_kinds=["paragraph", "table", "figure", "section"],
                              description="Remove a paragraph, table, figure or a whole section.",
                              params_schema={"type": "object", "additionalProperties": False,
                                             "properties": {}}),
            ])

    def created_kind(self, op: PlannedOperation) -> str:
        return {"add_heading": "section", "add_paragraph": "paragraph", "add_table": "table",
                "add_figure": "figure"}.get(op.op, "*")

    # ------------------------------------------------------------------ structure
    def _open(self, path: Path):
        docx = _docx()
        try:
            return docx.Document(str(path))
        except Exception as exc:
            raise AdapterError(f"cannot open document: {type(exc).__name__}: {exc}") from exc

    def _blocks(self, doc, root: str) -> tuple[list[dict[str, Any]], list[str]]:
        """Body elements in order with ids, kinds and section parents."""
        body = doc.element.body
        blocks = []
        issues = []
        seen: dict[str, int] = {}
        stack: list[tuple[int, str]] = []  # (level, section id)
        n_unreg = 0
        for i, el in enumerate(body.iterchildren()):
            tag = el.tag.split("}")[1]
            if tag not in ("p", "tbl"):
                continue
            marks = _bookmarks(el)
            ids = [m[3:] for m in marks]
            if len(ids) > 1:
                issues.append(f"element {i} carries several Daedelus bookmarks {ids}")
            cid = ids[0] if ids else None
            if tag == "tbl":
                kind = "table"
            else:
                lvl = _heading_level(el)
                kind = "section" if lvl is not None else (
                    "figure" if _has_picture(el) else "paragraph")
            if kind == "paragraph" and not _text(el).strip() and not cid:
                continue  # empty spacer paragraphs are not components
            if cid is None:
                n_unreg += 1
                cid = f"u{n_unreg}_{hashlib.sha256((_text(el) or str(i)).encode()).hexdigest()[:6]}"
                unreg = True
            else:
                unreg = False
            seen[cid] = seen.get(cid, 0) + 1
            if kind == "section":
                lvl = _heading_level(el)
                while stack and stack[-1][0] >= lvl:
                    stack.pop()
                parent = stack[-1][1] if stack else root
                stack.append((lvl, cid))
            else:
                parent = stack[-1][1] if stack else root
            blocks.append({"id": cid, "kind": kind, "el": el, "parent": parent,
                           "unregistered": unreg, "level": _heading_level(el) if tag == "p"
                           else None})
        for cid, n in seen.items():
            if n > 1:
                issues.append(f"component '{cid}' is ambiguous (bookmark used {n} times)")
        return blocks, issues

    def _bookmark(self, doc, el, cid: str) -> None:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        target = el
        if el.tag.endswith("}tbl"):
            target = el.find(f".//{_q('w:p')}")
        existing = [int(b.get(qn("w:id"))) for b in doc.element.body.iter(qn("w:bookmarkStart"))
                    if (b.get(qn("w:id")) or "").isdigit()]
        bid = str(max(existing, default=100) + 1)
        s = OxmlElement("w:bookmarkStart")
        s.set(qn("w:id"), bid)
        s.set(qn("w:name"), f"dd_{cid}")
        e = OxmlElement("w:bookmarkEnd")
        e.set(qn("w:id"), bid)
        ppr = target.find(qn("w:pPr"))
        idx = list(target).index(ppr) + 1 if ppr is not None else 0
        target.insert(idx, s)
        target.insert(idx + 1, e)

    # ------------------------------------------------------------------ lifecycle
    def create(self, native_dir: Path, template: str, params: dict[str, Any]) -> str:
        if template != "blank":
            raise AdapterError(f"unknown document template: {template}")
        docx = _docx()
        native_dir.mkdir(parents=True, exist_ok=True)
        doc = docx.Document()
        ids = {"version": 1, "root": params.get("root_id", "document"),
               "root_name": params.get("name", "Document"), "components": {}, "specs": {}}
        if params.get("title"):
            p = doc.add_heading(params["title"], level=0)
            self._bookmark(doc, p._p, "title")
        doc.save(str(native_dir / ENTRY))
        oc.save_ids(native_dir, ids)
        return ENTRY

    def _props(self, b) -> dict[str, Any]:
        el = b["el"]
        if b["kind"] == "table":
            rows = [[_text(tc) for tc in tr.iter(_q("w:tc"))] for tr in el.iter(_q("w:tr"))]
            return {"rows": rows[:100], "style": self._style(el)}
        st = self._style(el)
        return {"text": _text(el), "style": st, "level": b.get("level")}

    @staticmethod
    def _style(el) -> str | None:
        s = el.find(f".//{_q('w:pStyle')}") if el.tag.endswith("}p") else el.find(
            f".//{_q('w:tblStyle')}")
        return s.get(_q("w:val")) if s is not None else None

    def _state(self, doc, b) -> str:
        el = b["el"]
        if b["kind"] == "figure":
            blobs = []
            for blip in el.iter(_q("a:blip")):
                rid = blip.get(_q("r:embed"))
                try:
                    blobs.append(hashlib.sha256(doc.part.related_parts[rid].blob).hexdigest())
                except KeyError:
                    blobs.append(rid)
            ext = [x.get("cx") for x in el.iter("{http://schemas.openxmlformats.org/"
                                                 "drawingml/2006/wordprocessingDrawing}extent")]
            return oc.h({"img": blobs, "ext": ext, "text": _text(el)})
        xml = copy.deepcopy(el)
        for bm in list(xml.iter(_q("w:bookmarkStart"))) + list(xml.iter(_q("w:bookmarkEnd"))):
            bm.getparent().remove(bm)
        from lxml import etree

        return hashlib.sha256(etree.tostring(xml, method="c14n")).hexdigest()

    def inspect(self, native_dir: Path, entry: str) -> InspectResult:
        doc = self._open(native_dir / entry)
        ids = oc.load_ids(native_dir)
        root = ids.get("root", "document")
        blocks, issues = self._blocks(doc, root)
        comps = [Component(id=root, name=ids.get("root_name", "Document"), kind="document",
                           native_ref=".", metadata={"issues": issues})]
        states, meas, props = {}, {}, {}
        sec = doc.sections[0]
        states[root] = oc.h({"orient": str(sec.orientation), "margins": [
            sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin]})
        words = 0
        for b in blocks:
            t = _text(b["el"])
            label = (t[:48] or b["kind"]) if b["kind"] != "table" else f"Table {b['id']}"
            comps.append(Component(id=b["id"], name=label, kind=b["kind"], parent_id=b["parent"],
                                   native_ref=f"bookmark dd_{b['id']}" if not b["unregistered"]
                                   else "(unregistered)",
                                   metadata={"unregistered": b["unregistered"]}))
            states[b["id"]] = self._state(doc, b)
            props[b["id"]] = self._props(b)
            w = len(t.split())
            words += w
            if b["kind"] == "table":
                rows = props[b["id"]]["rows"]
                meas[b["id"]] = {"rows": len(rows), "columns": max((len(r) for r in rows),
                                                                    default=0), "words": w}
            else:
                meas[b["id"]] = {"words": w, "characters": len(t)}
        meas[root] = {"words": words, "paragraphs": sum(1 for b in blocks if b["kind"] ==
                                                         "paragraph"),
                      "figures": sum(1 for b in blocks if b["kind"] == "figure"),
                      "tables": sum(1 for b in blocks if b["kind"] == "table")}
        props[root] = {"orientation": "landscape" if sec.orientation == 1 else "portrait",
                       "issues": issues}
        return InspectResult(components=comps, states=states, measurements=meas,
                             properties=props, aggregates=[root])

    # ------------------------------------------------------------------ apply
    def apply(self, native_dir: Path, entry: str, operations: list[PlannedOperation],
              context: dict[str, Any]) -> ApplyResult:
        path = native_dir / entry
        bad = oc.unsupported_features(path, "docx")
        if bad:
            return ApplyResult(ok=False, error="editing refused - unsupported content: "
                               + ", ".join(bad) + ". An explicit, approved conversion is required.")
        doc = self._open(path)
        ids = oc.load_ids(native_dir)
        root = ids.get("root", "document")
        results = []
        for i, op in enumerate(operations):
            try:
                blocks, issues = self._blocks(doc, root)
                index = {b["id"]: b for b in blocks}
                amb = [x for x in issues if "ambiguous" in x]
                if op.component_id and op.component_id != root and \
                        self.op_spec(op.op).target_kinds != ["new"]:
                    if op.component_id not in index or any(f"'{op.component_id}'" in x
                                                           for x in amb):
                        raise AdapterError(f"component '{op.component_id}' not found or "
                                           f"ambiguous")
                detail = self._apply_one(doc, root, blocks, index, op, context)
                results.append({"index": i, "op": op.op, "component_id": op.component_id,
                                "status": "applied", "detail": detail})
            except Exception as exc:
                results.append({"index": i, "op": op.op, "component_id": op.component_id,
                                "status": "failed", "detail": f"{type(exc).__name__}: {exc}"})
                return ApplyResult(ok=False, results=results, error=results[-1]["detail"])
        # register every element that has no bookmark yet (stable ids from now on)
        blocks, _ = self._blocks(doc, root)
        for b in blocks:
            if b["unregistered"]:
                self._bookmark(doc, b["el"], b["id"])
        doc.save(str(path))
        oc.save_ids(native_dir, ids)
        return ApplyResult(ok=True, results=results)

    def _insert_point(self, doc, blocks, index, root, p) -> Any:
        """Element after which new content goes (None = end of body)."""
        after = p.get("after")
        parent = p.get("parent")
        if after:
            if after not in index:
                raise AdapterError(f"'after' component '{after}' not found")
            b = index[after]
            if b["kind"] == "section":
                return self._section_end(blocks, after)
            return b["el"]
        if parent and parent != root:
            if parent not in index:
                raise AdapterError(f"parent section '{parent}' not found")
            return self._section_end(blocks, parent)
        return None

    @staticmethod
    def _section_end(blocks, sid):
        ids = {sid}
        last = None
        for b in blocks:
            if b["id"] == sid:
                last = b["el"]
            elif b["parent"] in ids:
                ids.add(b["id"])
                last = b["el"]
        return last

    def _place(self, doc, new_el, after_el) -> None:
        body = doc.element.body
        if after_el is None:
            sect = body.find(_q("w:sectPr"))
            if sect is not None:
                sect.addprevious(new_el)
            else:
                body.append(new_el)
        else:
            after_el.addnext(new_el)

    def _image_path(self, key: str, context: dict[str, Any]) -> str:
        p = (context.get("files") or {}).get(key) or (context.get("source_paths") or {}).get(key)
        if not p or not Path(p).exists():
            raise AdapterError(f"image '{key}' is not available")
        return p

    def _apply_one(self, doc, root, blocks, index, op, context) -> str:
        from docx.shared import Cm

        p = op.params
        if op.op in ("add_heading", "add_paragraph", "add_table", "add_figure"):
            if p["id"] in index:
                return f"{p['id']} already exists (idempotent no-op)"
            after = self._insert_point(doc, blocks, index, root, p)
            if op.op == "add_heading":
                para = doc.add_heading(p["text"], level=int(p.get("level", 1)))
                el = para._p
            elif op.op == "add_paragraph":
                para = doc.add_paragraph(p["text"], style=p.get("style"))
                el = para._p
            elif op.op == "add_table":
                rows = p["rows"]
                ncol = max(len(r) for r in rows)
                tbl = doc.add_table(rows=len(rows), cols=ncol)
                tbl.style = p.get("style", "Light Grid Accent 1") if p.get("style") != "" else None
                self._fill_table(tbl, rows, p.get("header", True), p.get("number_format"))
                el = tbl._tbl
            else:
                para = doc.add_paragraph()
                para.add_run().add_picture(self._image_path(p["image"], context),
                                           width=Cm(float(p.get("width_cm", 15))))
                el = para._p
                if p.get("caption"):
                    cap = doc.add_paragraph(p["caption"], style="Caption")
                    self._place(doc, el, after)
                    self._bookmark(doc, el, p["id"])
                    el.addnext(cap._p)
                    self._bookmark(doc, cap._p, f"{p['id']}_caption")
                    return f"figure {p['id']} with caption"
            self._place(doc, el, after)
            self._bookmark(doc, el, p["id"])
            return f"{op.op.replace('add_', '')} {p['id']}"
        b = index.get(op.component_id) if op.component_id else None
        if op.op == "set_text":
            self._set_para_text(b["el"], p["text"])
            if p.get("style"):
                self._set_style(doc, b["el"], p["style"])
            return f"text of {op.component_id} set ({len(p['text'].split())} words)"
        if op.op == "apply_style":
            self._set_style(doc, b["el"], p["style"])
            return f"style {p['style']} on {op.component_id}"
        if op.op == "replace_section":
            kids = [x for x in blocks if x["parent"] == op.component_id and x["kind"] == "paragraph"]
            for k in kids:
                k["el"].getparent().remove(k["el"])
            anchor = b["el"]
            for j, text in enumerate(p["paragraphs"]):
                para = doc.add_paragraph(text, style=p.get("style"))
                anchor.addnext(para._p)
                self._bookmark(doc, para._p, f"{op.component_id}.p{j + 1}")
                anchor = para._p
            return f"section {op.component_id}: {len(kids)} paragraph(s) replaced by " \
                   f"{len(p['paragraphs'])}"
        if op.op == "update_table":
            from docx.table import Table

            tbl = Table(b["el"], doc._body)
            rows = p["rows"]
            while len(tbl.rows) < len(rows):
                tbl.add_row()
            while len(tbl.rows) > len(rows):
                tr = tbl.rows[-1]._tr
                tr.getparent().remove(tr)
            ncol = len(tbl.columns)
            if max(len(r) for r in rows) > ncol:
                raise AdapterError(f"table {op.component_id} has {ncol} columns; got "
                                   f"{max(len(r) for r in rows)}")
            self._fill_table(tbl, rows, None, p.get("number_format"))
            return f"table {op.component_id}: {len(rows)} row(s)"
        if op.op == "set_figure_image":
            blips = list(b["el"].iter(_q("a:blip")))
            if not blips:
                raise AdapterError(f"{op.component_id} has no image")
            path = self._image_path(p["image"], context)
            rid, _ = doc.part.get_or_add_image(path)
            blips[0].set(_q("r:embed"), rid)
            if p.get("width_cm"):
                from docx.shared import Cm as _Cm

                w = int(_Cm(float(p["width_cm"])))
                for ext in b["el"].iter("{http://schemas.openxmlformats.org/drawingml/2006/"
                                        "wordprocessingDrawing}extent"):
                    cx, cy = int(ext.get("cx")), int(ext.get("cy"))
                    ext.set("cx", str(w))
                    ext.set("cy", str(int(cy * w / cx)))
            return f"figure {op.component_id} image replaced"
        if op.op == "set_page_layout":
            from docx.enum.section import WD_ORIENT

            for sec in doc.sections:
                if p.get("orientation"):
                    want = WD_ORIENT.LANDSCAPE if p["orientation"] == "landscape" else \
                        WD_ORIENT.PORTRAIT
                    if sec.orientation != want:
                        sec.orientation = want
                        sec.page_width, sec.page_height = sec.page_height, sec.page_width
                if p.get("margins_cm"):
                    m = Cm(float(p["margins_cm"]))
                    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = m
            return f"page layout {p}"
        if op.op == "delete_component":
            if b["kind"] == "section":
                doomed = {b["id"]}
                for x in blocks:
                    if x["parent"] in doomed:
                        doomed.add(x["id"])
                for x in blocks:
                    if x["id"] in doomed:
                        x["el"].getparent().remove(x["el"])
                return f"removed section {b['id']} ({len(doomed)} element(s))"
            b["el"].getparent().remove(b["el"])
            return f"removed {b['kind']} {b['id']}"
        raise AdapterError(f"unknown operation {op.op}")

    @staticmethod
    def _fill_table(tbl, rows, header, number_format) -> None:
        for i, r in enumerate(rows):
            for j, v in enumerate(r):
                if isinstance(v, float) and number_format:
                    s = format(v, number_format)
                elif isinstance(v, float):
                    s = f"{v:,.2f}"
                else:
                    s = "" if v is None else str(v)
                cell = tbl.cell(i, j)
                para = cell.paragraphs[0]
                # keep paragraph properties (and bookmarks) - replace only the runs
                for rr in list(para._p.iter(_q("w:r"))):
                    rr.getparent().remove(rr)
                run = para.add_run(s)
                if header and i == 0:
                    run.bold = True

    @staticmethod
    def _set_para_text(p_el, text: str) -> None:
        runs = list(p_el.iter(_q("w:r")))
        if runs:
            first = runs[0]
            for r in runs[1:]:
                r.getparent().remove(r)
            for t in list(first.iter(_q("w:t"))):
                t.getparent().remove(t)
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn

            t = OxmlElement("w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = text
            first.append(t)
        else:
            from docx.text.paragraph import Paragraph

            Paragraph(p_el, None).add_run(text)

    @staticmethod
    def _set_style(doc, p_el, style: str) -> None:
        from docx.text.paragraph import Paragraph

        try:
            Paragraph(p_el, doc._body).style = doc.styles[style]
        except KeyError as exc:
            raise AdapterError(f"unknown style '{style}'") from exc

    # ------------------------------------------------------------------ preview etc.
    def structure(self, native_dir: Path, entry: str, out_dir: Path | None = None) -> dict[str, Any]:
        doc = self._open(native_dir / entry)
        ids = oc.load_ids(native_dir)
        blocks, issues = self._blocks(doc, ids.get("root", "document"))
        out = []
        for b in blocks:
            d = {"id": b["id"], "kind": b["kind"], "parent": b["parent"],
                 "unregistered": b["unregistered"], **self._props(b)}
            if b["kind"] == "figure" and out_dir is not None:
                for blip in b["el"].iter(_q("a:blip")):
                    part = doc.part.related_parts.get(blip.get(_q("r:embed")))
                    if part is not None:
                        ext = Path(part.partname).suffix or ".png"
                        f = out_dir / f"figure_{b['id']}{ext}"
                        f.write_bytes(part.blob)
                        d["image"] = f.name
                    break
            out.append(d)
        return {"blocks": out, "issues": issues}

    def preview(self, native_dir: Path, entry: str, out_dir: Path) -> dict[str, Path]:
        out_dir.mkdir(parents=True, exist_ok=True)
        st = out_dir / "structure.json"
        st.write_text(json.dumps(self.structure(native_dir, entry, out_dir), default=str))
        out: dict[str, Path] = {"structure": st}
        pdf = oc.lo_convert(native_dir / entry, "pdf", out_dir)
        if pdf is not None:
            out["pdf"] = pdf
            pngs = oc.pdf_to_pngs(pdf, out_dir, "page", max_pages=8, width=800)
            if pngs:
                out["render"] = pngs[0]
                for i, p in enumerate(pngs[1:], 2):
                    out[f"page{i}"] = p
        return out

    def export(self, native_dir: Path, entry: str, fmt: str, out_dir: Path) -> Path:
        out_dir.mkdir(parents=True, exist_ok=True)
        if fmt == "docx":
            dst = out_dir / entry
            shutil.copy2(native_dir / entry, dst)
            return dst
        if fmt == "pdf":
            pdf = oc.lo_convert(native_dir / entry, "pdf", out_dir)
            if pdf is None:
                raise AdapterError("PDF export needs LibreOffice")
            return pdf
        if fmt == "md":
            s = self.structure(native_dir, entry)
            lines = []
            for b in s["blocks"]:
                if b["kind"] == "section":
                    lines.append("#" * max(1, b.get("level") or 1) + " " + b["text"])
                elif b["kind"] == "table":
                    for i, r in enumerate(b["rows"]):
                        lines.append("| " + " | ".join(r) + " |")
                        if i == 0:
                            lines.append("|" + "---|" * len(r))
                else:
                    lines.append(b.get("text", ""))
                lines.append("")
            dst = out_dir / f"{Path(entry).stem}.md"
            dst.write_text("\n".join(lines))
            return dst
        raise AdapterError(f"document cannot export {fmt}")

    def diff(self, before_dir: Path, after_dir: Path, entry: str) -> str | None:
        try:
            a = self.structure(before_dir, entry)["blocks"]
            b = self.structure(after_dir, entry)["blocks"]
        except AdapterError:
            return None

        def lines(bl):
            out = []
            for x in bl:
                if x["kind"] == "table":
                    out += [f"[{x['id']}] | " + " | ".join(r) for r in x["rows"]]
                else:
                    out.append(f"[{x['id']}] {x.get('text', '')}")
            return out

        d = list(difflib.unified_diff(lines(a), lines(b), "before", "after", lineterm=""))
        return "\n".join(d) or None

    def validate(self, native_dir: Path, entry: str, checks: list[str],
                 context: dict[str, Any]) -> ValidationReport:
        rep = ValidationReport()
        try:
            doc = self._open(native_dir / entry)
            ids = oc.load_ids(native_dir)
            blocks, issues = self._blocks(doc, ids.get("root", "document"))
            rep.add("file_reopens", True, f"document reopened ({len(blocks)} components)")
            rep.add("component_identity", not issues, "; ".join(issues) or "all bookmarks unique")
        except AdapterError as exc:
            rep.add("file_reopens", False, str(exc))
            return rep
        if oc.soffice():
            import tempfile

            with tempfile.TemporaryDirectory() as td:
                ok = oc.lo_convert(native_dir / entry, "pdf", Path(td)) is not None
            rep.add("libreoffice_opens", ok, "LibreOffice rendered a copy to PDF" if ok else
                    "LibreOffice could not open a copy")
        return rep
