"""Presentation adapter (.pptx via python-pptx).

Components: presentation -> slides -> shapes (text, picture, autoshape, table, chart). Identity
uses PowerPoint's own ids - the slide id and the per-slide shape id are stored in the file
and are stable across editing. The sidecar only maps friendly Daedelus ids to them; shapes a
person adds in PowerPoint are addressable as ``<slide>.s<shape id>``. Everything is created
as editable slide objects (text frames, autoshapes, native tables and charts with embedded
data), never as flattened slide images.
"""

from __future__ import annotations

import copy
import hashlib
import json
import shutil
from pathlib import Path
from typing import Any

from ...models import Component, PlannedOperation, ValidationReport
from ..base import Adapter, AdapterError, AdapterInfo, ApplyResult, InspectResult, OperationSpec
from . import common as oc

ENTRY = "presentation.pptx"
ID = {"type": "string", "pattern": "[A-Za-z0-9_.-]+"}
HEX = {"type": "string", "pattern": "#[0-9a-fA-F]{6}"}
BOX = {"type": "array", "items": {"type": "number", "minimum": 0, "maximum": 1},
       "minItems": 4, "maxItems": 4, "description": "[x, y, w, h] as fractions of the slide"}
LAYOUTS = {"title": 0, "title_content": 1, "section": 2, "two_content": 3, "title_only": 5,
           "blank": 6}
TEXT_KINDS = ["text", "shape", "title"]


def _pptx():
    try:
        import pptx  # noqa: F401
    except ImportError as exc:  # pragma: no cover
        raise AdapterError("python-pptx is not installed") from exc
    import pptx

    return pptx


def _rgb(h: str):
    from pptx.dml.color import RGBColor

    return RGBColor.from_string(h.lstrip("#").upper())


def _shape_kind(sh) -> str:
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

    if sh.is_placeholder and sh.placeholder_format.type in (PP_PLACEHOLDER.TITLE,
                                                            PP_PLACEHOLDER.CENTER_TITLE):
        return "title"
    if getattr(sh, "has_chart", False) and sh.has_chart:
        return "chart"
    if getattr(sh, "has_table", False) and sh.has_table:
        return "table"
    if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return "image"
    if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
        return "shape"
    if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
        return "text"
    return "other"


class PresentationAdapter(Adapter):
    name = "presentation"
    version = "1"

    def check_environment(self) -> tuple[bool, str]:
        try:
            _pptx()
        except AdapterError as exc:
            return False, str(exc)
        return True, f"python-pptx; LibreOffice {'found' if oc.soffice() else 'not found (no thumbnails)'}"

    def info(self) -> AdapterInfo:
        ok, detail = self.check_environment()
        txt = {"text": {"type": "string"},
               "paragraphs": {"type": "array", "items": {"type": "string"}, "maxItems": 40},
               "size_pt": {"type": "number", "minimum": 6, "maximum": 120},
               "bold": {"type": "boolean"}, "color": HEX,
               "align": {"type": "string", "enum": ["left", "center", "right"]}}
        new_on_slide = {"id": ID, "slide": {"type": "string"}, "box": BOX}
        return AdapterInfo(
            name=self.name, version=self.version,
            description="PowerPoint presentations (.pptx) built from editable slide objects "
                        "with python-pptx; PowerPoint's slide and shape ids give stable "
                        "identities.",
            artifact_types=["presentation"], native_formats=[".pptx"],
            preview_formats=["png slide thumbnails (LibreOffice)", "structure (JSON)"],
            export_formats=["pptx", "pdf"],
            scopes=["presentation", "slide", "title", "text", "shape", "image", "table", "chart"],
            measurements=["slides", "shapes", "words", "series", "points", "sum"],
            validation=["file_reopens", "libreoffice_opens", "component_identity",
                        "component_preservation"],
            environment={"requires": "python-pptx (LibreOffice optional for thumbnails)"},
            available=ok, unavailable_reason=None if ok else detail, templates=["blank", "import"],
            operations=[
                OperationSpec(name="add_slide", family="create", aspects=["structure"],
                              target_kinds=["new"], description="Add a slide from a layout.",
                              params_schema={"type": "object", "required": ["id"],
                                             "additionalProperties": False, "properties": {
                                                 "id": ID, "layout": {"type": "string",
                                                                      "enum": list(LAYOUTS)},
                                                 "title": {"type": "string"},
                                                 "subtitle": {"type": "string"},
                                                 "index": {"type": "integer", "minimum": 0}}}),
                OperationSpec(name="move_slide", family="structure", aspects=["structure"],
                              target_kinds=["slide"], description="Move a slide to a position.",
                              params_schema={"type": "object", "required": ["index"],
                                             "additionalProperties": False, "properties": {
                                                 "index": {"type": "integer", "minimum": 0}}}),
                OperationSpec(name="set_text", family="content", aspects=["content", "text"],
                              target_kinds=TEXT_KINDS, description="Replace a text frame's text "
                              "(one paragraph per list item).",
                              params_schema={"type": "object", "additionalProperties": False,
                                             "properties": txt}),
                OperationSpec(name="add_textbox", family="create", aspects=["content", "text"],
                              target_kinds=["new"], description="Add a text box.",
                              params_schema={"type": "object", "required": ["id", "slide", "box"],
                                             "additionalProperties": False,
                                             "properties": {**new_on_slide, **txt}}),
                OperationSpec(name="add_shape", family="create", aspects=["style", "composition"],
                              target_kinds=["new"], description="Add an autoshape.",
                              params_schema={"type": "object", "required": ["id", "slide", "box"],
                                             "additionalProperties": False, "properties": {
                                                 **new_on_slide,
                                                 "shape": {"type": "string", "enum": [
                                                     "rectangle", "rounded_rectangle", "oval",
                                                     "chevron", "right_arrow"]},
                                                 "fill": HEX, "line": HEX, **txt}}),
                OperationSpec(name="add_image", family="create", aspects=["content", "figure"],
                              target_kinds=["new"], description="Add a picture ('image' is a "
                              "source id or engine-provided file key).",
                              params_schema={"type": "object",
                                             "required": ["id", "slide", "box", "image"],
                                             "additionalProperties": False, "properties": {
                                                 **new_on_slide, "image": {"type": "string"}}}),
                OperationSpec(name="set_image", family="figure", aspects=["figure"],
                              target_kinds=["image"], description="Replace a picture's image "
                              "(position and size kept).",
                              params_schema={"type": "object", "required": ["image"],
                                             "additionalProperties": False,
                                             "properties": {"image": {"type": "string"}}}),
                OperationSpec(name="add_table", family="create", aspects=["data", "content"],
                              target_kinds=["new"], description="Add a native table.",
                              params_schema={"type": "object",
                                             "required": ["id", "slide", "box", "rows"],
                                             "additionalProperties": False, "properties": {
                                                 **new_on_slide,
                                                 "rows": {"type": "array", "items": {"type": "array"}},
                                                 "size_pt": {"type": "number", "minimum": 6,
                                                             "maximum": 40}}}),
                OperationSpec(name="update_table", family="data", aspects=["data", "content"],
                              target_kinds=["table"], description="Replace a table's values "
                              "(same number of columns; rows are rebuilt as needed).",
                              params_schema={"type": "object", "required": ["rows"],
                                             "additionalProperties": False, "properties": {
                                                 "rows": {"type": "array", "items": {"type": "array"}}}}),
                OperationSpec(name="add_chart", family="create", aspects=["chart", "data"],
                              target_kinds=["new"], description="Add a native chart with embedded "
                              "data.",
                              params_schema={"type": "object",
                                             "required": ["id", "slide", "box", "categories",
                                                          "series"],
                                             "additionalProperties": False, "properties": {
                                                 **new_on_slide,
                                                 "type": {"type": "string", "enum": [
                                                     "column", "bar", "line", "pie"]},
                                                 "categories": {"type": "array"},
                                                 "series": {"type": "object"},
                                                 "title": {"type": "string"},
                                                 "color": HEX}}),
                OperationSpec(name="update_chart_data", family="data", aspects=["chart", "data"],
                              target_kinds=["chart"], description="Replace a chart's categories "
                              "and series (its embedded workbook is updated too).",
                              params_schema={"type": "object",
                                             "required": ["categories", "series"],
                                             "additionalProperties": False, "properties": {
                                                 "categories": {"type": "array"},
                                                 "series": {"type": "object"},
                                                 "title": {"type": "string"}}}),
                OperationSpec(name="set_shape_style", family="style", aspects=["style", "color"],
                              target_kinds=["shape", "text", "title"],
                              description="Fill, line and font colour of a shape.",
                              params_schema={"type": "object", "additionalProperties": False,
                                             "properties": {"fill": HEX, "line": HEX,
                                                            "font_color": HEX,
                                                            "size_pt": {"type": "number",
                                                                        "minimum": 6, "maximum": 120}}}),
                OperationSpec(name="apply_theme", family="style", aspects=["style", "color",
                                                                          "palette"],
                              target_kinds=["presentation", "slide"], subtree=True,
                              description="Consistent styling for slides: title/body fonts and "
                              "colours, accent colour, background.",
                              params_schema={"type": "object", "additionalProperties": False,
                                             "properties": {
                                                 "title_font": {"type": "string"},
                                                 "body_font": {"type": "string"},
                                                 "title_color": HEX, "body_color": HEX,
                                                 "background": HEX, "accent": HEX}}),
                OperationSpec(name="delete_component", family="structure", aspects=["structure"],
                              target_kinds=["slide", "title", "text", "shape", "image", "table",
                                            "chart"],
                              description="Remove a slide or a shape.",
                              params_schema={"type": "object", "additionalProperties": False,
                                             "properties": {}}),
            ])

    def created_kind(self, op: PlannedOperation) -> str:
        return {"add_slide": "slide", "add_textbox": "text", "add_shape": "shape",
                "add_image": "image", "add_table": "table", "add_chart": "chart"}.get(op.op, "*")

    # ------------------------------------------------------------------ identity
    def _open(self, path: Path):
        pptx = _pptx()
        try:
            return pptx.Presentation(str(path))
        except Exception as exc:
            raise AdapterError(f"cannot open presentation: {type(exc).__name__}: {exc}") from exc

    def _index(self, prs, ids: dict[str, Any]):
        """component id -> (kind, slide, shape|None); plus issues."""
        alias = {(v["slide_id"], v.get("shape_id")): k for k, v in ids["components"].items()}
        out: dict[str, Any] = {}
        issues = []
        for s in prs.slides:
            sid = alias.get((s.slide_id, None)) or f"slide{s.slide_id}"
            out[sid] = ("slide", s, None)
            for sh in s.shapes:
                cid = alias.get((s.slide_id, sh.shape_id)) or f"{sid}.s{sh.shape_id}"
                if cid in out:
                    issues.append(f"component '{cid}' is ambiguous")
                out[cid] = (_shape_kind(sh), s, sh)
        present = {(v[1].slide_id, v[2].shape_id if v[2] is not None else None)
                   for v in out.values()}
        for k, v in ids["components"].items():
            if (v["slide_id"], v.get("shape_id")) not in present:
                issues.append(f"component '{k}' is missing (slide {v['slide_id']}, shape "
                              f"{v.get('shape_id')})")
        return out, issues

    # ------------------------------------------------------------------ lifecycle
    def create(self, native_dir: Path, template: str, params: dict[str, Any]) -> str:
        if template == "import":
            try:
                oc.import_copy(params["path"], native_dir, ENTRY, params, "presentation",
                               "Presentation", (".pptx",))
            except (KeyError, ValueError) as exc:
                raise AdapterError(f"import: {exc}") from exc
            _pptx().Presentation(str(native_dir / ENTRY))  # must parse
            return ENTRY
        if template != "blank":
            raise AdapterError(f"unknown presentation template: {template}")
        pptx = _pptx()
        native_dir.mkdir(parents=True, exist_ok=True)
        prs = pptx.Presentation()  # default template layouts (4:3) keep placeholders aligned
        prs.save(str(native_dir / ENTRY))
        oc.save_ids(native_dir, {"version": 1, "root": params.get("root_id", "presentation"),
                                 "root_name": params.get("name", "Presentation"),
                                 "components": {}, "specs": {}})
        return ENTRY

    def _xml_hash(self, el) -> str:
        from lxml import etree

        return hashlib.sha256(etree.tostring(copy.deepcopy(el), method="c14n")).hexdigest()

    def _shape_state(self, sh) -> str:
        parts = [self._xml_hash(sh._element)]
        if getattr(sh, "has_chart", False) and sh.has_chart:
            parts.append(self._xml_hash(sh.chart._chartSpace))
        if _shape_kind(sh) == "image":
            parts.append(hashlib.sha256(sh.image.blob).hexdigest())
        return oc.h(parts)

    def _props(self, prs, kind, sh) -> dict[str, Any]:
        W, H = prs.slide_width, prs.slide_height
        d: dict[str, Any] = {"name": sh.name, "box": [round(sh.left / W, 4), round(sh.top / H, 4),
                                                      round(sh.width / W, 4),
                                                      round(sh.height / H, 4)]
                             if sh.width is not None else None}
        if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
            d["text"] = sh.text_frame.text
        if kind == "table":
            d["rows"] = [[c.text for c in r.cells] for r in sh.table.rows]
        if kind == "chart":
            ch = sh.chart
            d["chart_type"] = str(ch.chart_type)
            d["categories"] = [str(c) for c in ch.plots[0].categories] if ch.plots else []
            d["series"] = {s.name: list(s.values) for s in ch.plots[0].series} if ch.plots else {}
            d["title"] = ch.chart_title.text_frame.text if ch.has_title else None
        return d

    def inspect(self, native_dir: Path, entry: str) -> InspectResult:
        prs = self._open(native_dir / entry)
        ids = oc.load_ids(native_dir)
        root = ids.get("root", "presentation")
        idx, issues = self._index(prs, ids)
        comps = [Component(id=root, name=ids.get("root_name", "Presentation"),
                           kind="presentation", native_ref=".", metadata={"issues": issues})]
        states = {root: oc.h({"order": [s.slide_id for s in prs.slides],
                              "theme": ids.get("specs", {}).get("theme")})}
        meas: dict[str, dict[str, Any]] = {root: {"slides": len(prs.slides)}}
        props: dict[str, dict[str, Any]] = {root: {"issues": issues,
                                                   "theme": ids.get("specs", {}).get("theme"),
                                                   "slide_order": []}}
        slide_cid = {}
        for cid, (kind, s, sh) in idx.items():
            if kind == "slide":
                slide_cid[s.slide_id] = cid
        for n, s in enumerate(prs.slides):
            cid = slide_cid[s.slide_id]
            title = s.shapes.title.text_frame.text if s.shapes.title is not None else ""
            comps.append(Component(id=cid, name=f"{n + 1}. {title or s.slide_layout.name}",
                                   kind="slide", parent_id=root, native_ref=f"slide id {s.slide_id}",
                                   metadata={"index": n}))
            bg = s.background.fill
            states[cid] = oc.h({"layout": s.slide_layout.name,
                                "bg": str(bg.fore_color.rgb) if bg.type == 1 else None})
            words = 0
            for cid2, (kind, s2, sh) in idx.items():
                if sh is None or s2.slide_id != s.slide_id:
                    continue
                pr = self._props(prs, kind, sh)
                label = (pr.get("text") or pr.get("title") or sh.name)[:40]
                comps.append(Component(id=cid2, name=label, kind=kind, parent_id=cid,
                                       native_ref=f"slide {s.slide_id} shape {sh.shape_id}"))
                states[cid2] = self._shape_state(sh)
                props[cid2] = pr
                w = len((pr.get("text") or "").split())
                words += w
                m: dict[str, Any] = {"words": w}
                if kind == "chart":
                    vals = [v for vs in pr["series"].values() for v in vs if v is not None]
                    m.update(series=len(pr["series"]), points=len(pr["categories"]),
                             sum=round(sum(vals), 6) if vals else None)
                if kind == "table":
                    m.update(rows=len(pr["rows"]), columns=len(pr["rows"][0]) if pr["rows"] else 0)
                meas[cid2] = m
            meas[cid] = {"shapes": len(s.shapes), "words": words, "index": n}
            props[cid] = {"index": n, "layout": s.slide_layout.name, "title": title}
            props[root]["slide_order"].append(cid)
        return InspectResult(components=comps, states=states, measurements=meas,
                             properties=props, aggregates=[root])

    # ------------------------------------------------------------------ apply
    def apply(self, native_dir: Path, entry: str, operations: list[PlannedOperation],
              context: dict[str, Any]) -> ApplyResult:
        path = native_dir / entry
        own = self._chart_embeddings(path)  # chart data workbooks are part of native charts
        bad = [b for b in oc.unsupported_features(path, "pptx", own)]
        if bad:
            return ApplyResult(ok=False, error="editing refused - unsupported content: "
                               + ", ".join(bad) + ". An explicit, approved conversion is required.")
        prs = self._open(path)
        ids = oc.load_ids(native_dir)
        results = []
        for i, op in enumerate(operations):
            try:
                idx, issues = self._index(prs, ids)
                spec = self.op_spec(op.op)
                if op.component_id and op.component_id != ids.get("root", "presentation") and \
                        spec.target_kinds != ["new"]:
                    if op.component_id not in idx or any(f"'{op.component_id}'" in x
                                                         for x in issues):
                        raise AdapterError(f"component '{op.component_id}' not found or ambiguous")
                detail = self._apply_one(prs, ids, idx, op, context)
                results.append({"index": i, "op": op.op, "component_id": op.component_id,
                                "status": "applied", "detail": detail})
            except Exception as exc:
                results.append({"index": i, "op": op.op, "component_id": op.component_id,
                                "status": "failed", "detail": f"{type(exc).__name__}: {exc}"})
                return ApplyResult(ok=False, results=results, error=results[-1]["detail"])
        prs.save(str(path))
        oc.save_ids(native_dir, ids)
        return ApplyResult(ok=True, results=results)

    @staticmethod
    def _chart_embeddings(path: Path) -> set[str]:
        import posixpath
        import re
        import zipfile

        own = set()
        with zipfile.ZipFile(path) as z:
            for n in z.namelist():
                if n.startswith("ppt/charts/_rels/") and n.endswith(".rels"):
                    for t in re.findall(r'Target="([^"]+)"', z.read(n).decode()):
                        own.add(posixpath.normpath(posixpath.join("ppt/charts", t)))
        return own

    def _slide(self, idx, key: str):
        if key not in idx or idx[key][0] != "slide":
            raise AdapterError(f"slide '{key}' not found")
        return idx[key][1]

    def _box(self, prs, box):
        x, y, w, h = box
        return (int(x * prs.slide_width), int(y * prs.slide_height), int(w * prs.slide_width),
                int(h * prs.slide_height))

    def _image(self, key: str, context: dict[str, Any]) -> str:
        p = (context.get("files") or {}).get(key) or (context.get("source_paths") or {}).get(key)
        if not p or not Path(p).exists():
            raise AdapterError(f"image '{key}' is not available")
        return p

    def _register(self, ids, cid, slide, shape=None) -> None:
        ids["components"][cid] = {"slide_id": slide.slide_id,
                                  **({"shape_id": shape.shape_id} if shape is not None else {})}

    def _write_text(self, tf, p: dict[str, Any], theme: dict[str, Any] | None, title: bool):
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Pt

        paras = p.get("paragraphs") if p.get("paragraphs") is not None else \
            [p.get("text", "")]
        tf.clear()
        for j, text in enumerate(paras):
            para = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            run = para.add_run()
            run.text = str(text)
            f = run.font
            th = theme or {}
            font = th.get("title_font" if title else "body_font")
            if font:
                f.name = font
            color = p.get("color") or th.get("title_color" if title else "body_color")
            if color:
                f.color.rgb = _rgb(color)
            if p.get("size_pt"):
                f.size = Pt(float(p["size_pt"]))
            if p.get("bold") is not None:
                f.bold = bool(p["bold"])
            if p.get("align"):
                para.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                                  "right": PP_ALIGN.RIGHT}[p["align"]]

    def _apply_one(self, prs, ids, idx, op: PlannedOperation, context) -> str:
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Pt

        p = op.params
        theme = ids.get("specs", {}).get("theme")
        if op.op == "add_slide":
            if p["id"] in ids["components"]:
                return f"slide {p['id']} already exists (idempotent no-op)"
            layout = prs.slide_layouts[LAYOUTS[p.get("layout", "title_content")]]
            s = prs.slides.add_slide(layout)
            if p.get("title") is not None and s.shapes.title is not None:
                self._write_text(s.shapes.title.text_frame, {"text": p["title"]}, theme, True)
            if p.get("subtitle") is not None and len(s.placeholders) > 1:
                self._write_text(s.placeholders[1].text_frame, {"text": p["subtitle"]}, theme,
                                 False)
            if p.get("index") is not None:
                self._move(prs, s, int(p["index"]))
            if theme and theme.get("background"):
                s.background.fill.solid()
                s.background.fill.fore_color.rgb = _rgb(theme["background"])
            self._register(ids, p["id"], s)
            if s.shapes.title is not None:
                self._register(ids, f"{p['id']}.title", s, s.shapes.title)
            if len(s.placeholders) > 1:
                self._register(ids, f"{p['id']}.body", s, s.placeholders[1])
            return f"slide {p['id']} ({p.get('layout', 'title_content')})"
        if op.op == "move_slide":
            s = idx[op.component_id][1]
            self._move(prs, s, int(p["index"]))
            return f"slide {op.component_id} -> position {p['index']}"
        if op.op == "set_text":
            kind, s, sh = idx[op.component_id]
            if not sh.has_text_frame:
                raise AdapterError(f"{op.component_id} has no text frame")
            self._write_text(sh.text_frame, p, theme, kind == "title")
            return f"text of {op.component_id} set"
        if op.op in ("add_textbox", "add_shape", "add_image", "add_table", "add_chart"):
            if p["id"] in ids["components"]:
                return f"{p['id']} already exists (idempotent no-op)"
            s = self._slide(idx, p["slide"])
            x, y, w, h = self._box(prs, p["box"])
            if op.op == "add_textbox":
                sh = s.shapes.add_textbox(x, y, w, h)
                sh.text_frame.word_wrap = True
                self._write_text(sh.text_frame, p, theme, False)
            elif op.op == "add_shape":
                kind = {"rectangle": MSO_SHAPE.RECTANGLE,
                        "rounded_rectangle": MSO_SHAPE.ROUNDED_RECTANGLE,
                        "oval": MSO_SHAPE.OVAL, "chevron": MSO_SHAPE.CHEVRON,
                        "right_arrow": MSO_SHAPE.RIGHT_ARROW}[p.get("shape", "rectangle")]
                sh = s.shapes.add_shape(kind, x, y, w, h)
                fill = p.get("fill") or (theme or {}).get("accent")
                if fill:
                    sh.fill.solid()
                    sh.fill.fore_color.rgb = _rgb(fill)
                if p.get("line"):
                    sh.line.color.rgb = _rgb(p["line"])
                if p.get("text") or p.get("paragraphs"):
                    self._write_text(sh.text_frame, p, theme, False)
            elif op.op == "add_image":
                sh = s.shapes.add_picture(self._image(p["image"], context), x, y, w, h)
            elif op.op == "add_table":
                rows = p["rows"]
                ncol = max(len(r) for r in rows)
                gs = s.shapes.add_table(len(rows), ncol, x, y, w, h)
                sh = gs
                self._fill_table(gs.table, rows, p.get("size_pt"))
            else:
                ctype = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED,
                         "bar": XL_CHART_TYPE.BAR_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS,
                         "pie": XL_CHART_TYPE.PIE}[p.get("type", "column")]
                cd = self._chart_data(p["categories"], p["series"])
                sh = s.shapes.add_chart(ctype, x, y, w, h, cd)
                ch = sh.chart
                if p.get("title"):
                    ch.has_title = True
                    ch.chart_title.text_frame.text = p["title"]
                color = p.get("color") or (theme or {}).get("accent")
                if color and p.get("type", "column") != "pie":
                    for ser in ch.plots[0].series:
                        fmt = ser.format
                        if p.get("type") == "line":
                            fmt.line.color.rgb = _rgb(color)
                        else:
                            fmt.fill.solid()
                            fmt.fill.fore_color.rgb = _rgb(color)
                ch.has_legend = len(p["series"]) > 1
            sh.name = p["id"]
            self._register(ids, p["id"], s, sh)
            return f"{op.op.replace('add_', '')} {p['id']} on {p['slide']}"
        if op.op == "set_image":
            kind, s, sh = idx[op.component_id]
            from pptx.parts.image import Image as PImage

            path = self._image(p["image"], context)
            img_part, rid = s.part.get_or_add_image_part(path)
            sh._element.blipFill.blip.rEmbed = rid
            _ = PImage
            return f"image of {op.component_id} replaced"
        if op.op == "update_table":
            kind, s, sh = idx[op.component_id]
            tbl = sh.table
            rows = p["rows"]
            ncol = len(tbl.columns)
            if max(len(r) for r in rows) != ncol:
                raise AdapterError(f"table {op.component_id} has {ncol} columns")
            if len(rows) != len(tbl.rows):
                # rebuild in place: same box and name, same Daedelus id (new native shape id)
                x, y, w, h = sh.left, sh.top, sh.width, sh.height
                el = sh._element
                el.getparent().remove(el)
                gs = s.shapes.add_table(len(rows), ncol, x, y, w, h)
                gs.name = sh.name
                self._fill_table(gs.table, rows, None)
                self._register(ids, op.component_id, s, gs)
                return f"table {op.component_id} rebuilt with {len(rows)} row(s)"
            self._fill_table(tbl, rows, None)
            return f"table {op.component_id} updated"
        if op.op == "update_chart_data":
            kind, s, sh = idx[op.component_id]
            sh.chart.replace_data(self._chart_data(p["categories"], p["series"]))
            if p.get("title"):
                sh.chart.has_title = True
                sh.chart.chart_title.text_frame.text = p["title"]
            return f"chart {op.component_id}: {len(p['categories'])} categories, " \
                   f"{len(p['series'])} series"
        if op.op == "set_shape_style":
            kind, s, sh = idx[op.component_id]
            if p.get("fill"):
                sh.fill.solid()
                sh.fill.fore_color.rgb = _rgb(p["fill"])
            if p.get("line"):
                sh.line.color.rgb = _rgb(p["line"])
            if (p.get("font_color") or p.get("size_pt")) and sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    for r in para.runs:
                        if p.get("font_color"):
                            r.font.color.rgb = _rgb(p["font_color"])
                        if p.get("size_pt"):
                            r.font.size = Pt(float(p["size_pt"]))
            return f"style on {op.component_id}"
        if op.op == "apply_theme":
            theme = {**(ids.get("specs", {}).get("theme") or {}), **p}
            ids.setdefault("specs", {})["theme"] = theme
            root = ids.get("root", "presentation")
            slides = [s for s in prs.slides] if op.component_id in (None, root) else \
                [idx[op.component_id][1]]
            for s in slides:
                if theme.get("background"):
                    s.background.fill.solid()
                    s.background.fill.fore_color.rgb = _rgb(theme["background"])
                for sh in s.shapes:
                    if not getattr(sh, "has_text_frame", False) or not sh.has_text_frame:
                        continue
                    is_title = _shape_kind(sh) == "title"
                    for para in sh.text_frame.paragraphs:
                        for r in para.runs:
                            font = theme.get("title_font" if is_title else "body_font")
                            col = theme.get("title_color" if is_title else "body_color")
                            if font:
                                r.font.name = font
                            if col:
                                r.font.color.rgb = _rgb(col)
            return f"theme applied to {len(slides)} slide(s)"
        if op.op == "delete_component":
            kind, s, sh = idx[op.component_id]
            if kind == "slide":
                lst = prs.slides._sldIdLst
                for sld in list(lst):
                    if int(sld.get("id")) == s.slide_id:
                        prs.part.drop_rel(sld.rId)
                        lst.remove(sld)
                for k in [k for k, v in ids["components"].items() if v["slide_id"] == s.slide_id]:
                    ids["components"].pop(k)
                return f"removed slide {op.component_id}"
            sh._element.getparent().remove(sh._element)
            ids["components"].pop(op.component_id, None)
            return f"removed {kind} {op.component_id}"
        raise AdapterError(f"unknown operation {op.op}")

    @staticmethod
    def _chart_data(categories, series):
        from pptx.chart.data import CategoryChartData

        cd = CategoryChartData()
        cd.categories = [str(c) for c in categories]
        for name, vals in series.items():
            if len(vals) != len(categories):
                raise AdapterError(f"series '{name}' has {len(vals)} values for "
                                   f"{len(categories)} categories")
            cd.add_series(str(name), [None if v is None else float(v) for v in vals])
        return cd

    @staticmethod
    def _fill_table(tbl, rows, size_pt):
        from pptx.util import Pt

        for i, r in enumerate(rows):
            for j, v in enumerate(r):
                cell = tbl.cell(i, j)
                cell.text = f"{v:,.2f}" if isinstance(v, float) else ("" if v is None else str(v))
                if size_pt:
                    for para in cell.text_frame.paragraphs:
                        for run in para.runs:
                            run.font.size = Pt(float(size_pt))

    @staticmethod
    def _move(prs, slide, index: int) -> None:
        lst = prs.slides._sldIdLst
        items = list(lst)
        el = next(x for x in items if int(x.get("id")) == slide.slide_id)
        lst.remove(el)
        lst.insert(min(index, len(items) - 1), el)

    # ------------------------------------------------------------------ preview etc.
    def structure(self, native_dir: Path, entry: str) -> dict[str, Any]:
        insp = self.inspect(native_dir, entry)
        root = next(c for c in insp.components if c.parent_id is None)
        slides = []
        for sid in insp.properties[root.id]["slide_order"]:
            shapes = [{"id": c.id, "kind": c.kind, **insp.properties.get(c.id, {})}
                      for c in insp.components if c.parent_id == sid]
            slides.append({"id": sid, **insp.properties[sid], "shapes": shapes})
        return {"slides": slides, "theme": insp.properties[root.id].get("theme"),
                "issues": insp.properties[root.id].get("issues")}

    def preview(self, native_dir: Path, entry: str, out_dir: Path) -> dict[str, Path]:
        out_dir.mkdir(parents=True, exist_ok=True)
        st = out_dir / "structure.json"
        st.write_text(json.dumps(self.structure(native_dir, entry), default=str), encoding="utf-8")
        out: dict[str, Path] = {"structure": st}
        pdf = oc.lo_convert(native_dir / entry, "pdf", out_dir)
        if pdf is not None:
            out["pdf"] = pdf
            pngs = oc.pdf_to_pngs(pdf, out_dir, "slide", max_pages=30, width=960)
            if pngs:
                out["render"] = pngs[0]
                for i, p in enumerate(pngs[1:], 2):
                    out[f"slide{i}"] = p
        return out

    def export(self, native_dir: Path, entry: str, fmt: str, out_dir: Path) -> Path:
        out_dir.mkdir(parents=True, exist_ok=True)
        if fmt == "pptx":
            dst = out_dir / entry
            shutil.copy2(native_dir / entry, dst)
            return dst
        if fmt == "pdf":
            pdf = oc.lo_convert(native_dir / entry, "pdf", out_dir)
            if pdf is None:
                raise AdapterError("PDF export needs LibreOffice")
            return pdf
        raise AdapterError(f"presentation cannot export {fmt}")

    def diff(self, before_dir: Path, after_dir: Path, entry: str) -> str | None:
        try:
            a, b = self.structure(before_dir, entry), self.structure(after_dir, entry)
        except AdapterError:
            return None

        def lines(st):
            out = []
            for n, s in enumerate(st["slides"], 1):
                out.append(f"slide {n} [{s['id']}] {s.get('title') or ''}")
                for sh in s["shapes"]:
                    t = sh.get("text") or sh.get("title") or ""
                    extra = f" rows={sh['rows']}" if sh.get("rows") else ""
                    extra += f" series={sh['series']}" if sh.get("series") else ""
                    out.append(f"  {sh['kind']} [{sh['id']}] {t[:80]}{extra[:200]}")
            return out

        import difflib

        d = list(difflib.unified_diff(lines(a), lines(b), "before", "after", lineterm=""))
        return "\n".join(d) or None

    def validate(self, native_dir: Path, entry: str, checks: list[str],
                 context: dict[str, Any]) -> ValidationReport:
        rep = ValidationReport()
        try:
            prs = self._open(native_dir / entry)
            idx, issues = self._index(prs, oc.load_ids(native_dir))
            rep.add("file_reopens", True, f"presentation reopened ({len(prs.slides)} slide(s))")
            rep.add("component_identity", not issues, "; ".join(issues) or "all ids resolved")
            flat = [k for k, v in idx.items() if v[0] == "image" and len(idx) > 0]
            editable = sum(1 for v in idx.values() if v[0] in ("text", "title", "shape",
                                                               "table", "chart"))
            rep.add("editable_objects", editable > 0 or not len(prs.slides),
                    f"{editable} editable object(s), {len(flat)} picture(s)")
        except AdapterError as exc:
            rep.add("file_reopens", False, str(exc))
            return rep
        if oc.soffice():
            import tempfile

            with tempfile.TemporaryDirectory() as td:
                ok = oc.lo_convert(native_dir / entry, "pdf", Path(td)) is not None
            rep.add("libreoffice_opens", ok, "LibreOffice rendered a copy to PDF" if ok else
                    "LibreOffice could not open a copy")
        return rep
